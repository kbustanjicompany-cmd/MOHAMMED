"""Areas of the closed polylines drawn by the user in 3-3.dxf (layer 0: yellow = area 1,
blue = area 2), and excavation under them with the agreed levels:
  area 1: 1114.85 -> 1113.0 (1.85 m), the two F.B.L 1109.5 rafts 1114.85 -> 1109.5 (5.35 m)
  area 2: 1120.9 -> 1118.0 (2.90 m)
Output: attached_areas.xlsx
"""
import ezdxf
from ezdxf import path as zpath
from shapely.geometry import Polygon
import openpyxl
from openpyxl.styles import Font, PatternFill

SRC = "3-3.dxf"
NAMES = {  # handle -> (element, zone, from, to)
    "421E": ("Raft/mat 60 cm - north row + east strip + south-east block", 1, 1114.85, 1113.0),
    "422B": ("Raft 60 cm - south-west block", 1, 1114.85, 1113.0),
    "422F": ("Raft 60 cm - south-west strip", 1, 1114.85, 1113.0),
    "422C": ("Raft 80 cm - middle strip", 1, 1114.85, 1113.0),
    "422E": ("Raft 80 cm - west strip (F.B.L 1109.5)", 1, 1114.85, 1109.5),
    "422D": ("Raft 50 cm - north-west (F.B.L 1109.5)", 1, 1114.85, 1109.5),
    "4221": ("Footing F4", 1, 1114.85, 1113.0), "4220": ("Footing F5", 1, 1114.85, 1113.0),
    "421F": ("Footing F5", 1, 1114.85, 1113.0), "4222": ("Footing F6", 1, 1114.85, 1113.0),
    "4223": ("Footing F7", 1, 1114.85, 1113.0), "4224": ("Footing F7", 1, 1114.85, 1113.0),
    "4227": ("Footing F1", 1, 1114.85, 1113.0), "4228": ("Footing F1", 1, 1114.85, 1113.0),
    "4226": ("Footing F8", 1, 1114.85, 1113.0), "422A": ("Footing F8", 1, 1114.85, 1113.0),
    "4225": ("Footing F9", 1, 1114.85, 1113.0), "4229": ("Footing F10", 1, 1114.85, 1113.0),
    "4230": ("Area 2 - foundations (blue outline)", 2, 1120.9, 1118.0),
}
msp = ezdxf.readfile(SRC).modelspace()
polys = {}
for e in msp.query('LWPOLYLINE[layer=="0"]'):
    if e.dxf.handle in NAMES:
        polys[e.dxf.handle] = Polygon([(v.x, v.y) for v in zpath.make_path(e).flattening(0.005)])

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Areas"
ws.append(["Handle", "Element", "Area", "Area (m²)", "From", "To", "Depth (m)", "Excavation (m³)"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
tot = {1: [0, 0], 2: [0, 0]}
for h, (el, z, t, b) in NAMES.items():
    a = polys[h].area
    ws.append([h, el, f"Area {z}", round(a, 2), t, b, round(t - b, 2), round(a * (t - b), 2)])
    tot[z][0] += a; tot[z][1] += a * (t - b)
ws.append([])
for z in (1, 2):
    ws.append(["", f"TOTAL area {z}", "", round(tot[z][0], 2), "", "", "", round(tot[z][1], 2)])
ws.append(["", "TOTAL", "", round(tot[1][0] + tot[2][0], 2), "", "", "", round(tot[1][1] + tot[2][1], 2)])
for r in range(ws.max_row - 2, ws.max_row + 1):
    for c in ws[r]: c.font = Font(bold=True)
for col, w in zip("ABCDEFGH", [8, 52, 9, 12, 9, 9, 10, 15]):
    ws.column_dimensions[col].width = w
wb.save("attached_areas.xlsx")
for h, (el, z, t, b) in NAMES.items():
    print(h, el, z, round(polys[h].area, 2), round(polys[h].area * (t - b), 2))
print({z: [round(v, 2) for v in s] for z, s in tot.items()}, round(tot[1][0] + tot[2][0], 2), round(tot[1][1] + tot[2][1], 2))
