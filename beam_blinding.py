"""Blinding under the ground beams only (G.B, G.B1-G.B5), 100 mm projection each side.

  - beam concrete = the foundation RC footprint minus the footings/rafts (parts narrower
    than 1 m, see footings_concrete.py); pieces under 0.3 m² (corner slivers) are ignored
  - beam blinding = beam concrete grown by 100 mm, minus the blinding already under the
    footings and rafts (so nothing is counted twice)
Outputs: beam_blinding.dxf (traced closed polylines, plan coordinates), beam_blinding.pdf,
beam_blinding.xlsx
"""
import sys
import numpy as np
from skimage import measure
from shapely.geometry import Polygon
from shapely.ops import unary_union
import ezdxf
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.argv = sys.argv[:1]
exec(open("footings_concrete.py").read().split('ws.title = "Footings & rafts"')[0])
beam_rc = rc & ~feet
lb, nb = ndi.label(beam_rc)
sizes = ndi.sum(np.ones_like(lb), lb, index=np.arange(1, nb + 1)) * RES * RES
beam_rc = np.isin(lb, [i + 1 for i, a in enumerate(sizes) if a > 0.3])
P = int(round(PROJECTION / RES)); pdisk = np.hypot(*np.mgrid[-P:P + 1, -P:P + 1]) <= P
beam_bl = ndi.binary_dilation(beam_rc, structure=pdisk) & ~feet_blind
conc_area = beam_rc.sum() * RES * RES
bl_area = beam_bl.sum() * RES * RES
drawn_area = (blind & ~feet_blind).sum() * RES * RES


def trace(mask):
    m = np.pad(mask, 1)
    geom = Polygon()
    rings = []
    for c in measure.find_contours(m.astype(np.uint8), 0.5):
        if len(c) >= 4:
            pg = Polygon([(X0 + (cc - 1) * RES + RES / 2, Y1 - (rr - 1) * RES - RES / 2) for rr, cc in c]).buffer(0)
            if pg.area > 0.02:
                rings.append(pg)
    for pg in sorted(rings, key=lambda g: -g.area):
        geom = geom.symmetric_difference(pg)
    return geom.simplify(0.01)


g = trace(beam_bl)
polys = [p for p in getattr(g, "geoms", [g]) if p.area > 0.05]
doc = ezdxf.new("R2018", setup=True); doc.units = ezdxf.units.M
msp = doc.modelspace()
doc.layers.add("BLINDING-GROUND-BEAMS", color=30); doc.layers.add("BLINDING-GROUND-BEAMS-HATCH", color=30)
doc.layers.add("BLINDING-LABEL", color=7)
for pg in polys:
    ext = list(pg.exterior.coords)[:-1]
    msp.add_lwpolyline(ext, close=True, dxfattribs={"layer": "BLINDING-GROUND-BEAMS"})
    h = msp.add_hatch(color=30, dxfattribs={"layer": "BLINDING-GROUND-BEAMS-HATCH"})
    h.set_pattern_fill("ANSI31", scale=0.03)
    h.paths.add_polyline_path(ext, is_closed=True, flags=ezdxf.const.BOUNDARY_PATH_EXTERNAL)
    for hole in pg.interiors:
        msp.add_lwpolyline(list(hole.coords)[:-1], close=True, dxfattribs={"layer": "BLINDING-GROUND-BEAMS"})
        h.paths.add_polyline_path(list(hole.coords)[:-1], is_closed=True)
msp.add_mtext(f"BLINDING UNDER GROUND BEAMS (100 mm projection each side)\\P"
              f"beam concrete = {conc_area:.2f} m²\\Pbeam blinding = {bl_area:.2f} m²  (volume @ 100 mm = {bl_area * .1:.2f} m³)",
              dxfattribs={"layer": "BLINDING-LABEL", "char_height": 0.5}).set_location((-50, 66), attachment_point=7)
doc.saveas("beam_blinding.dxf")

fig, ax = plt.subplots(figsize=(16.5, 13))
xs = np.linspace(X0, X1, W); ys = np.linspace(Y1, Y0, H)
ax.contourf(xs, ys, feet_blind.astype(float), levels=[.5, 1.5], colors=["#dddddd"])
for pg in polys:
    ax.fill(*pg.exterior.xy, fc="#e69138", ec="#7f4000", lw=.4)
ax.set_xlim(X0, X1); ax.set_ylim(Y0, Y1); ax.set_aspect("equal"); ax.axis("off")
ax.fill([], [], fc="#e69138", label=f"ground beam blinding: {bl_area:,.2f} m²")
ax.fill([], [], fc="#dddddd", label="footings / rafts blinding (not included)")
ax.legend(loc="upper center", bbox_to_anchor=(.5, -.01), ncol=2, fontsize=12, frameon=False)
ax.set_title(f"Blinding under ground beams only — 100 mm each side — {bl_area:,.2f} m²", fontsize=14, weight="bold")
fig.savefig("beam_blinding.pdf", bbox_inches="tight")

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ground beams"
ws.append(["Item", "Value"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for row in [("Ground beam concrete footprint (m²)", round(conc_area, 2)),
            ("Ground beam blinding, 100 mm each side, outside footing/raft blinding (m²)", round(bl_area, 2)),
            ("Blinding volume @ 100 mm (m³)", round(bl_area * .1, 2)),
            ("Equivalent length @ 0.50 m blinding width (m)", round(bl_area / .5, 1)),
            ("Check: beam blinding as drawn on the plan (m²)", round(drawn_area, 2))]:
    ws.append(list(row))
ws.column_dimensions["A"].width = 70; ws.column_dimensions["B"].width = 14
wb.save("beam_blinding.xlsx")
print(f"beam concrete={conc_area:.2f} beam blinding={bl_area:.2f} drawn={drawn_area:.2f} pieces={len(polys)}")
