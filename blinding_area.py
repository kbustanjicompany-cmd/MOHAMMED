"""Blinding (lean concrete) area under the foundations, from FOUNDATIONS.dxf.

Method
  - The foundation plan's blinding outlines are the S-BLINDING LINE layer (lines,
    polylines, arcs). They are drawn on a 2 cm raster.
  - The lines are thickened by 0.3 m so that small breaks in the outlines (up to
    ~0.6 m) do not let regions leak into each other; each enclosed region is then
    grown back by the same 0.3 m, so the areas are measured to the red line.
  - An enclosed region is blinding when the concrete outline (S-CONCRETE PLANS) runs
    just inside its edge (blinding projects ~10 cm beyond the concrete); otherwise it
    is an empty pocket between beams. Regions holding a RAFT tag or the water
    collecting tank are always blinding.
  - One open end of the west raft strip is closed by hand (see GAP_FIXES).
Outputs: blinding_area.pdf/png, blinding_area.xlsx
"""
import ezdxf, numpy as np
from ezdxf import path as zpath
from scipy import ndimage as ndi
from PIL import Image, ImageDraw
from shapely.geometry import LineString, box
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill

THICKNESS = 0.10                       # "100mm CONCRETE BLINDING" (drawing notes)
RES, GAP = 0.02, 0.30                  # raster size, gap bridging (m)
X0, Y0, X1, Y1 = -50, -8, 30, 60       # foundation plan window (drawing units = m)
GAP_FIXES = [((-33.56, 18.95), (-33.17, 17.75))]
FORCE_BLINDING = [(-40.63, 34.33),     # RAFT Thickness 80cm (west strip)
                  (22.0, 14.0)]        # water collecting tank (section A-A)

msp = ezdxf.readfile("FOUNDATIONS.dxf").modelspace()
plan = box(X0, Y0, X1, Y1)


def lines(layer):
    out = []
    for e in msp.query(f'*[layer=="{layer}"]'):
        ents = e.virtual_entities() if e.dxftype() == "INSERT" else [e]
        for v in ents:
            if v.dxftype() not in ("LINE", "LWPOLYLINE", "ARC", "POLYLINE", "CIRCLE"):
                continue
            pts = [(p.x, p.y) for p in zpath.make_path(v).flattening(0.01)]
            if len(pts) > 1 and LineString(pts).intersects(plan):
                out.append(LineString(pts))
    return out


blind_lines = lines("S-BLINDING LINE") + [LineString(g) for g in GAP_FIXES]
conc_lines = lines("S-CONCRETE PLANS")
W, H = int((X1 - X0) / RES), int((Y1 - Y0) / RES)
to_px = lambda x, y: ((x - X0) / RES, (Y1 - y) / RES)


def raster(gs):
    im = Image.new("1", (W, H), 0); dr = ImageDraw.Draw(im)
    for g in gs:
        dr.line([to_px(x, y) for x, y in g.coords], fill=1, width=1)
    return np.array(im, bool)


Lb, Lc = raster(blind_lines), raster(conc_lines)
R = int(GAP / RES)
disk = np.hypot(*np.mgrid[-R:R + 1, -R:R + 1]) <= R
lab, n = ndi.label(~ndi.binary_dilation(Lb, structure=disk))
outside = set(np.unique(np.r_[lab[0], lab[-1], lab[:, 0], lab[:, -1]])) - {0}
forced = {lab[int(to_px(x, y)[1]), int(to_px(x, y)[0])] for x, y in FORCE_BLINDING}
band = int(0.12 / RES)
pockets = []
for k, sl in enumerate(ndi.find_objects(lab), 1):
    if k in outside or k in forced:
        continue
    pad = R + band + 2
    s0 = slice(max(sl[0].start - pad, 0), min(sl[0].stop + pad, H))
    s1 = slice(max(sl[1].start - pad, 0), min(sl[1].stop + pad, W))
    reg = ndi.binary_dilation(lab[s0, s1] == k, structure=disk)
    edge = reg & ~ndi.binary_erosion(reg, iterations=2)
    inner = reg & ~ndi.binary_erosion(reg, iterations=band) & ~edge
    if (Lc[s0, s1] & inner).sum() / max(edge.sum() / 2, 1) < 0.35:
        pockets.append(k)
blind = ~ndi.binary_dilation(np.isin(lab, list(outside) + pockets), structure=disk)
total = blind.sum() * RES * RES

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Blinding"
ws.append(["Item", "Value"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
ws.append(["Blinding area (m²)", round(total, 2)])
ws.append([f"Blinding thickness (m)", THICKNESS])
ws.append(["Blinding volume (m³)", round(total * THICKNESS, 2)])
ws.append(["Method", "area enclosed by S-BLINDING LINE in the foundation plan (footings, ground beams, rafts)"])
for col, w in zip("AB", [24, 90]):
    ws.column_dimensions[col].width = w
wb.save("blinding_area.xlsx")

fig, ax = plt.subplots(figsize=(16, 14), dpi=110)
xs = np.linspace(X0, X1, W); ys = np.linspace(Y1, Y0, H)
ax.contourf(xs, ys, blind.astype(float), levels=[0.5, 1.5], colors=["#f6b26b"], alpha=0.7)
for g in conc_lines: ax.plot(*g.xy, color="#00a0c0", lw=0.4)
for g in blind_lines: ax.plot(*g.xy, color="red", lw=0.7)
ax.set_xlim(X0, X1); ax.set_ylim(Y0, Y1); ax.set_aspect("equal"); ax.axis("off")
ax.set_title(f"Blinding under foundations (footings, ground beams, rafts)\n"
             f"Total area = {total:,.2f} m²   —   volume @ {THICKNESS * 100:g} cm = {total * THICKNESS:,.2f} m³",
             fontsize=14, weight="bold")
fig.savefig("blinding_area.png", bbox_inches="tight"); fig.savefig("blinding_area.pdf", bbox_inches="tight")
print(f"blinding area={total:.2f} m2 volume={total * THICKNESS:.2f} m3")
