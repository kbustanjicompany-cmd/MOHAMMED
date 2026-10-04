"""Cyclopean concrete (rubble + concrete) under the hatched footings of area 2 (syclopien.dxf).

- Footings = grey SOLID hatches on layer "S-D-WATER PROOF" (drawn on the blinding outline):
  the retaining-wall footing (east strip, south-east curve, south strip and the north end)
  with the pads attached to it, plus three separate pads.
- Cyclopean outline = footings + OFFSET (1.0 m) all round, EXCEPT on the outer face of the
  retaining-wall footing (soil / boundary side): there it stops flush with the footing edge.
- Cyclopean thickness DEPTH = 0.5 m, under the footing founding level (area 2: 1118.0 -> 1117.5).
- Extra excavation because of this item (footing pits were priced over the blinding outline
  from 1120.9 down to 1118.0, see footing_excavation.py):
    1) the cyclopean layer itself: cyclopean area x 0.5
    2) widening the pit by the offset band from 1120.9 down to 1118.0:
       (cyclopean area - footing area) x 2.9
- Extra backfill: the widened band above the cyclopean, 1118.0 -> 1120.9 = item 2.
Outputs: cyclopean.xlsx, cyclopean.pdf, cyclopean.dxf
"""
import numpy as np
import ezdxf
from ezdxf import path
from shapely.geometry import Polygon, LineString
from shapely.ops import unary_union
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill

SRC = "syclopien.dxf"
OFFSET, DEPTH = 1.0, 0.5
TOP, FBL = 1120.9, 1118.0          # area 2 bulk excavation level and footing founding level

doc = ezdxf.readfile(SRC); msp = doc.modelspace()
pieces = [Polygon([(v.x, v.y) for v in p.flattening(0.02)]).buffer(0)
          for h in msp.query('HATCH[layer=="S-D-WATER PROOF"]') for p in path.from_hatch(h)]
foot = unary_union(pieces)
ret = [p for p in pieces if p.area > 30]               # retaining-wall footing pieces (180 + 39 m²)
pads = [p for p in pieces if p.area <= 30]

# outer face of the retaining-wall footing: walk its outline from the north-west end, along the
# north edge, down the east edge, round the outer arc and along the south edge to the west end
north, south = sorted(ret, key=lambda p: -p.bounds[3])
nc, sc = list(north.exterior.coords), list(south.exterior.coords)
outer = [nc[0], nc[1], nc[2]]                          # north edge + east edge (down to the curve)
k0 = max(range(len(sc)), key=lambda i: sc[i][0])       # outer arc starts at the east-most point
arc = sc[k0:-1]                                        # outer arc + south edge + west end pad
outer += arc
# extend both ends far into the site so the half-plane cut covers the whole offset band
def ext(p, q, L=200):
    d = np.subtract(p, q); d = d / np.hypot(*d); return tuple(np.add(p, d * L))
start = ext(outer[0], outer[1]); end = ext(outer[-1], outer[-2])
inside = Polygon([start] + outer + [end, (end[0] - 300, end[1] + 400), (start[0] - 300, start[1] + 300)]).buffer(0)
assert inside.buffer(0.01).contains(foot), "footings must lie on the site side of the outer face"

cyc = unary_union([p.buffer(OFFSET, join_style=2) for p in pieces]).intersection(inside)
cyc = unary_union([cyc, foot])
A_f, A_c = foot.area, cyc.area
band = A_c - A_f
V_cyc = A_c * DEPTH
V_ex1, V_ex2 = V_cyc, band * (TOP - FBL)
V_exc, V_fill = V_ex1 + V_ex2, band * (TOP - FBL)

# check: offset measured from the concrete face (hatch is the blinding, 0.10 m bigger)
cyc_c = unary_union([unary_union([p.buffer(OFFSET - 0.1, join_style=2) for p in pieces]).intersection(inside), foot])

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Cyclopean"
hdr = lambda: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row]]
ws.append(["Item", "Area (m²)", "From level", "To level", "Depth (m)", "Volume (m³)"]); hdr()
ws.append(["Hatched footings (blinding outline)", round(A_f, 2)])
ws.append([f"Cyclopean concrete (footings + {OFFSET} m offset, flush on the retaining-wall outer face)",
           round(A_c, 2), FBL, FBL - DEPTH, DEPTH, round(V_cyc, 2)])
ws.append([])
ws.append(["EXTRA EXCAVATION", "", "", "", "", ""]); ws[ws.max_row][0].font = Font(bold=True)
ws.append(["1) Cyclopean layer under the footings", round(A_c, 2), FBL, FBL - DEPTH, DEPTH, round(V_ex1, 2)])
ws.append([f"2) Widening the pits by the offset band", round(band, 2), TOP, FBL, TOP - FBL, round(V_ex2, 2)])
ws.append(["Total extra excavation", "", "", "", "", round(V_exc, 2)]); ws[ws.max_row][0].font = Font(bold=True)
ws.append([])
ws.append(["EXTRA BACKFILL", "", "", "", "", ""]); ws[ws.max_row][0].font = Font(bold=True)
ws.append(["Offset band above the cyclopean", round(band, 2), FBL, TOP, TOP - FBL, round(V_fill, 2)])
ws.append([])
ws.append(["Check - offset 1 m from the concrete face (0.9 m from the hatch)", round(cyc_c.area, 2), "", "", DEPTH,
           round(cyc_c.area * DEPTH, 2)])
ws2 = wb.create_sheet("Footings")
ws2.append(["No.", "Footing", "Area (m²)"])
for i, p in enumerate(sorted(pieces, key=lambda p: -p.area), 1):
    ws2.append([i, "Retaining-wall footing" if p in ret else "Pad footing", round(p.area, 2)])
ws.column_dimensions["A"].width = 70
for col in "BCDEF": ws.column_dimensions[col].width = 12
wb.save("cyclopean.xlsx")

# ---- detailed table per footing: the cyclopean outline is shared out to the nearest footing
from shapely import points as _pts, contains_xy as _cxy, distance as _dist
names = []
_ri = _pi = 0
for p in sorted(pieces, key=lambda p: (p not in ret, -p.bounds[3])):
    if p in ret:
        _ri += 1; names.append((p, f"Retaining-wall footing - {'north/east' if _ri == 1 else 'south/curve'} part"))
    else:
        _pi += 1; names.append((p, f"Pad footing P{_pi}"))
STEP = 0.05
def share(region):
    x0, y0, x1, y1 = region.bounds
    gx, gy = np.meshgrid(np.arange(x0 + STEP / 2, x1, STEP), np.arange(y0 + STEP / 2, y1, STEP))
    gx, gy = gx.ravel(), gy.ravel(); m = _cxy(region, gx, gy); gx, gy = gx[m], gy[m]
    pt = _pts(gx, gy)
    d = np.vstack([_dist(p, pt) for p, _ in names]); own = d.argmin(axis=0)
    a = np.bincount(own, minlength=len(names)) * STEP * STEP
    return a * region.area / a.sum()                  # scale out the grid error
H = TOP - FBL
detail = []
for cyc_r in (cyc, cyc_c):
    ac = share(cyc_r)
    rows_ = []
    for (p, n), a in zip(names, ac):
        b = a - p.area
        rows_.append(dict(name=n, af=p.area, ac=a, band=b, exc0=p.area * H, exc1=a * H, dexc=b * H,
                          lay=a * DEPTH, cyc=a * DEPTH, fill=b * H, tot=b * H + a * DEPTH))
    detail.append(rows_)
for title, rows_ in (("Detail (offset from hatch)", detail[0]), ("Detail (offset from concrete)", detail[1])):
    wd = wb.create_sheet(title)
    wd.append([f"Area 2 hatched footings - excavation {TOP} -> {FBL} ({H:.1f} m), cyclopean {DEPTH} m "
               f"({FBL} -> {FBL - DEPTH}); offset {'1.0 m from the hatch (blinding) outline' if 'hatch' in title else '1.0 m from the concrete face (0.9 m from the hatch)'}; "
               "no offset on the retaining-wall outer face"])
    wd.append([])
    wd.append(["Footing", "Blinding area (m²)", "Cyclopean area (m²)", "Offset band (m²)",
               f"Excavation before - blinding x {H:.1f} (m³)", f"Excavation after - cyclopean x {H:.1f} (m³)",
               "Difference to 1118 (m³)", f"Cyclopean layer excavation x {DEPTH} (m³)", "Total extra excavation (m³)",
               "Cyclopean concrete (m³)", f"Extra backfill - band x {H:.1f} (m³)"])
    for c in wd[wd.max_row]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="center")
    keys = ["af", "ac", "band", "exc0", "exc1", "dexc", "lay", "tot", "cyc", "fill"]
    for r in rows_:
        wd.append([r["name"]] + [round(r[k], 2) for k in keys])
    wd.append(["TOTAL"] + [round(sum(r[k] for r in rows_), 2) for k in keys])
    for c in wd[wd.max_row]: c.font = Font(bold=True)
    wd.column_dimensions["A"].width = 38
    for col in "BCDEFGHIJK": wd.column_dimensions[col].width = 15
    wd.row_dimensions[3].height = 62
    wd.freeze_panes = "B4"
wb.move_sheet("Detail (offset from hatch)", offset=-2); wb.move_sheet("Detail (offset from concrete)", offset=-2)
wb.active = 0
wb.save("cyclopean.xlsx")
for rows_ in detail:
    for r in rows_: print(r["name"], *[round(r[k], 2) for k in ("af", "ac", "band", "exc0", "exc1", "dexc", "lay", "tot", "fill")])
    print("TOTAL", *[round(sum(r[k] for r in rows_), 2) for k in ("af", "ac", "band", "exc0", "exc1", "dexc", "lay", "tot", "fill")])

out = ezdxf.new("R2018"); om = out.modelspace()
out.layers.add("FOOTINGS-HATCHED", color=8); out.layers.add("CYCLOPEAN", color=1)
for g in getattr(cyc, "geoms", [cyc]):
    om.add_lwpolyline(list(g.exterior.coords)[:-1], close=True, dxfattribs={"layer": "CYCLOPEAN"})
for p in pieces:
    om.add_lwpolyline(list(p.exterior.coords)[:-1], close=True, dxfattribs={"layer": "FOOTINGS-HATCHED"})
out.saveas("cyclopean.dxf")

fig, ax = plt.subplots(figsize=(9, 16))
for e in msp.query("LINE LWPOLYLINE ARC"):
    if e.dxf.layer in ("S-CONCRETE PLANS", "S-S-BEARING WALL UNDER THE GROUND", "S-S-BEARING WALL"):
        try: pts = np.array([(v.x, v.y) for v in path.make_path(e).flattening(0.05)])
        except Exception: continue
        ax.plot(pts[:, 0], pts[:, 1], color="#bbbbbb", lw=0.5, zorder=1)
for g in getattr(cyc, "geoms", [cyc]):
    ax.fill(*g.exterior.xy, fc="#f4cccc", ec="#c00000", lw=1.5, zorder=2)
for p in pieces:
    ax.fill(*p.exterior.xy, fc="#999999", ec="black", lw=0.8, zorder=3)
ax.plot(*np.array(outer).T, color="#00a000", lw=3, zorder=4, label="retaining-wall outer face (no offset)")
ax.legend(loc="upper left", bbox_to_anchor=(0, -0.02))
ax.set_xlim(foot.bounds[0] - 4, foot.bounds[2] + 3); ax.set_ylim(foot.bounds[1] - 3, foot.bounds[3] + 3)
ax.set_aspect("equal"); ax.axis("off")
ax.set_title(f"Cyclopean concrete under hatched footings (area 2)\n"
             f"area {A_c:.2f} m² x {DEPTH} m = {V_cyc:.2f} m³\n"
             f"extra excavation {V_exc:.2f} m³ — extra backfill {V_fill:.2f} m³", fontsize=12)
_k = 0
for p in sorted(pads, key=lambda p: -p.bounds[3]):
    _k += 1; ax.text(*p.centroid.coords[0], f"P{_k}", ha="center", va="center", fontsize=11, weight="bold", color="white", zorder=5)
fig.savefig("cyclopean.pdf", bbox_inches="tight"); fig.savefig("cyclopean.png", dpi=80, bbox_inches="tight")
print(f"footings {A_f:.2f}  cyclopean {A_c:.2f} m² -> {V_cyc:.2f} m³ ; band {band:.2f} ; "
      f"extra exc {V_ex1:.2f}+{V_ex2:.2f}={V_exc:.2f} ; fill {V_fill:.2f} ; check(0.9) {cyc_c.area:.2f}")
