"""Arabic, formatted version of the cyclopean quantities (cyclopean_ar.xlsx).
Areas come from cyclopean.py (DXF geometry); all volumes are live Excel formulas."""
import io, contextlib
with contextlib.redirect_stdout(io.StringIO()):
    exec(open("cyclopean.py").read())
ROWS = detail                       # per-footing rows: [offset from concrete (adopted), offset from hatch]
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

AR = {"Retaining-wall footing - north/east part": "قاعدة الجدار الاستنادي – الجزء الشمالي والشرقي",
      "Retaining-wall footing - south/curve part": "قاعدة الجدار الاستنادي – الجزء الجنوبي والمنحنى"}
ar_name = lambda n: AR.get(n, n.replace("Pad footing", "قاعدة منفصلة"))

F = "Arial"
NAVY, LIGHT, BAND, GOLD = "1F3864", "D9E2F3", "F2F2F2", "FFF2CC"
thin = Side(style="thin", color="A6A6A6"); med = Side(style="medium", color=NAVY)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0.00;-#,##0.00;"-"'
LVL = '0.00'
center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
right = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)

wb = openpyxl.Workbook()
S = wb.active; S.title = "الملخص"
SUM = "'الملخص'"

def setup(ws, widths):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

def title(ws, rng, text, sub=None):
    a = rng.split(":")[0]
    ws.merge_cells(rng); c = ws[a]; c.value = text
    c.font = Font(name=F, size=15, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = center
    ws.row_dimensions[c.row].height = 32
    if sub:
        r = c.row + 1; rng2 = rng.replace(str(c.row), str(r))
        ws.merge_cells(rng2); s = ws[rng2.split(":")[0]]; s.value = sub
        s.font = Font(name=F, size=10, italic=True, color="404040"); s.alignment = center
        ws.row_dimensions[r].height = 30

def head(ws, row, labels, col0=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col0 + i, t)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = center; c.border = box
    ws.row_dimensions[row].height = 48

def cell(ws, r, c, v, fmt=None, bold=False, fill=None, color="000000", align=center):
    x = ws.cell(r, c, v); x.font = Font(name=F, size=10, bold=bold, color=color); x.border = box; x.alignment = align
    if fmt: x.number_format = fmt
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    return x

# ---------------- summary sheet: parameters
setup(S, {"A": 4, "B": 46, "C": 16, "D": 16, "E": 16, "F": 4})
title(S, "B2:E2", "كميات الخرسانة السيكلوبية (دبش + خرسانة) – القواعد المهشرة – المنطقة الثانية",
      "الحفر والردم الإضافيان الناتجان عن بند السيكلوبين مقارنةً بالحفر على مساحة النظافة")
r = 5
S.merge_cells(f"B{r}:E{r}"); c = S[f"B{r}"]; c.value = "أولاً: المعطيات"
c.font = Font(name=F, size=12, bold=True, color=NAVY); c.alignment = right
r += 1
head(S, r, ["البند", "القيمة", "الوحدة", "ملاحظة"], 2); S.row_dimensions[r].height = 24
params = [("منسوب الحفر العام (بداية حفر القواعد)", TOP, "م", "مُدخل"),
          ("منسوب الحفر النهائي = منسوب أسفل السيكلوبين", FBL, "م", "مُدخل"),
          ("سماكة طبقة السيكلوبين", DEPTH, "م", "مُدخل"),
          ("منسوب سطح السيكلوبين = منسوب تأسيس القواعد", "=C8+C9", "م", "محسوب"),
          ("عمق الحفر من المنسوب العام إلى منسوب الحفر النهائي", "=C7-C8", "م", "محسوب"),
          ("ارتفاع الردم فوق السيكلوبين حتى المنسوب العام", "=C7-C10", "م", "محسوب"),
          ("إزاحة السيكلوبين حول القواعد", OFFSET, "م", "لا إزاحة على الوجه الخارجي للجدار الاستنادي")]
for i, (t, v, u, n) in enumerate(params):
    rr = r + 1 + i
    cell(S, rr, 2, t, align=right, fill=BAND if i % 2 else None)
    inp = n == "مُدخل"
    cell(S, rr, 3, v, LVL, bold=True, fill=GOLD if inp else (BAND if i % 2 else None), color="0000FF" if inp else "000000")
    cell(S, rr, 4, u, fill=BAND if i % 2 else None)
    cell(S, rr, 5, n, fill=BAND if i % 2 else None).font = Font(name=F, size=9, color="595959")
S["C13"].comment = Comment("للعلم فقط: المساحات في صفحات التفصيل مقيسة من الرسم بهذه الإزاحة؛ تغيير هذه الخلية لا يغيّر المساحات.", "Claude")
P_TOP, P_FBL, P_T, P_CT, P_H, P_HF = (f"{SUM}!$C${k}" for k in (7, 8, 9, 10, 11, 12))

# ---------------- detail sheets
def detail(name, cyc_rows, note):
    ws = wb.create_sheet(name)
    setup(ws, {"A": 2, "B": 5, "C": 40, **{k: 14 for k in "DEFGHIJK"}, "L": 2})
    title(ws, "B2:K2", f"جدول تفصيلي – {name}", note)
    hr = 5
    head(ws, hr, ["م", "القاعدة", "مساحة النظافة\n(م²)", "مساحة السيكلوبين\n(م²)", "مساحة شريط الإزاحة\n(م²)",
                  "الحفر قبل\nعلى مساحة النظافة\n(م³)", "الحفر بعد\nعلى مساحة السيكلوبين\n(م³)", "الحفر الزائد\n(م³)",
                  "خرسانة السيكلوبين\n(م³)", "الردم الزائد\n(م³)"], 2)
    sub = hr + 1
    for i, t in enumerate(["", "", "من الرسم", "من الرسم", "سيكلوبين − نظافة", "نظافة × عمق الحفر",
                           "سيكلوبين × عمق الحفر", "بعد − قبل", "سيكلوبين × السماكة", "شريط × ارتفاع الردم"]):
        c = cell(ws, sub, 2 + i, t, fill=LIGHT); c.font = Font(name=F, size=8, italic=True, color=NAVY)
    r0 = sub + 1
    for i, rw in enumerate(cyc_rows):
        r = r0 + i; f = BAND if i % 2 else None
        cell(ws, r, 2, i + 1, fill=f)
        cell(ws, r, 3, ar_name(rw["name"]), align=right, fill=f)
        cell(ws, r, 4, round(rw["af"], 2), NUM, fill=f, color="0000FF")
        cell(ws, r, 5, round(rw["ac"], 2), NUM, fill=f, color="0000FF")
        cell(ws, r, 6, f"=E{r}-D{r}", NUM, fill=f)
        cell(ws, r, 7, f"=D{r}*{P_H}", NUM, fill=f)
        cell(ws, r, 8, f"=E{r}*{P_H}", NUM, fill=f)
        cell(ws, r, 9, f"=H{r}-G{r}", NUM, fill=f, bold=True)
        cell(ws, r, 10, f"=E{r}*{P_T}", NUM, fill=f, bold=True)
        cell(ws, r, 11, f"=F{r}*{P_HF}", NUM, fill=f, bold=True)
    rt = r0 + len(cyc_rows)
    ws.merge_cells(start_row=rt, start_column=2, end_row=rt, end_column=3)
    cell(ws, rt, 2, "المجموع", bold=True, fill=LIGHT); cell(ws, rt, 3, None, fill=LIGHT)
    for col in "DEFGHIJK":
        c = cell(ws, rt, " BCDEFGHIJK".index(col) + 1, f"=SUM({col}{r0}:{col}{rt - 1})", NUM, bold=True, fill=LIGHT)
        c.border = Border(left=thin, right=thin, top=med, bottom=med)
    ws.cell(rt, 2).border = ws.cell(rt, 3).border = Border(left=thin, right=thin, top=med, bottom=med)
    n = rt + 2
    notes = ["ملاحظات:",
             "• مساحات النظافة والسيكلوبين مقيسة من ملف syclopien.dxf (التهشير الرمادي على طبقة S-D-WATER PROOF، وهو مرسوم على حد النظافة).",
             "• في أماكن تداخل إزاحات القواعد المتجاورة وُزّعت المساحة المشتركة على أقرب قاعدة حتى لا تُحسب مرتين.",
             "• الحفر الزائد = توسيع حفر القواعد بشريط الإزاحة من المنسوب العام حتى منسوب الحفر النهائي؛ لا يوجد حفر تحت هذا المنسوب.",
             "• الردم الزائد = شريط الإزاحة فوق السيكلوبين حتى المنسوب العام. لا يشمل الردم فوق القواعد نفسها لأنه موجود في الحالتين.",
             "• الجوانب محسوبة رأسية؛ أي ميول في جوانب الحفر تزيد الحفر والردم.",
             "• الخلايا الزرقاء قيم مقيسة من الرسم، والباقي معادلات مرتبطة بصفحة الملخص."]
    for i, t in enumerate(notes):
        ws.merge_cells(start_row=n + i, start_column=2, end_row=n + i, end_column=11)
        c = ws.cell(n + i, 2, t); c.alignment = right
        c.font = Font(name=F, size=10 if i else 11, bold=not i, color=NAVY if not i else "404040")
    ws.freeze_panes = ws.cell(r0, 4)
    ws.print_title_rows = f"{hr}:{sub}"
    return name, rt

S1 = detail("الإزاحة 1 م من وجه الخرسانة", ROWS[0],
            "المعتمد: الإزاحة 1.0 م مقيسة من وجه الخرسانة (0.9 م من حد النظافة) – بدون إزاحة على الوجه الخارجي للجدار الاستنادي")
S2 = detail("بديل - الإزاحة من حد النظافة", ROWS[1],
            "للمقارنة فقط: الإزاحة 1.0 م مقيسة من خط التهشير (حد النظافة) – بدون إزاحة على الوجه الخارجي للجدار الاستنادي")

# ---------------- summary table
r = 16
S.merge_cells(f"B{r}:E{r}"); c = S[f"B{r}"]; c.value = "ثانياً: ملخص الكميات"
c.font = Font(name=F, size=12, bold=True, color=NAVY); c.alignment = right
r += 1
head(S, r, ["البند", "المعتمد: الإزاحة\nمن وجه الخرسانة", "للمقارنة: الإزاحة\nمن حد النظافة", "الوحدة"], 2)
items = [("مساحة النظافة (القواعد المهشرة)", "D", "م²"), ("مساحة السيكلوبين", "E", "م²"),
         ("مساحة شريط الإزاحة", "F", "م²"), ("الحفر قبل – على مساحة النظافة", "G", "م³"),
         ("الحفر بعد – على مساحة السيكلوبين", "H", "م³"), ("الحفر الزائد", "I", "م³"),
         ("خرسانة السيكلوبين (دبش + خرسانة)", "J", "م³"), ("الردم الزائد", "K", "م³")]
key = {"I", "J", "K"}
for i, (t, col, u) in enumerate(items):
    rr = r + 1 + i; f = GOLD if col in key else (BAND if i % 2 else None)
    cell(S, rr, 2, t, align=right, fill=f, bold=col in key)
    for j, (nm, rt) in enumerate((S1, S2)):
        cell(S, rr, 3 + j, f"='{nm}'!{col}{rt}", NUM, fill=f, bold=col in key, color="008000")
    cell(S, rr, 5, u, fill=f)
n = r + len(items) + 2
for i, t in enumerate(["• المعتمد: إزاحة 1 م من وجه الخرسانة (العمود الأول)؛ العمود الثاني للمقارنة فقط.",
                       "• الخلايا الصفراء في المعطيات قابلة للتعديل وتُحدّث كل الجداول تلقائياً."]):
    S.merge_cells(f"B{n + i}:E{n + i}"); c = S[f"B{n + i}"]; c.value = t
    c.font = Font(name=F, size=10, color="404040"); c.alignment = right
wb.save("cyclopean_ar.xlsx")

# ---- cached values: LibreOffice is not usable here, so evaluate the (simple) formulas in Python
# and write them into the file as cached results; Excel still recalculates them normally.
import re, zipfile, shutil
from openpyxl.utils import range_boundaries, get_column_letter
REF = re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
cache = {}
def val(sh, coord):
    k = (sh, coord)
    if k in cache: return cache[k]
    v = wb[sh][coord].value
    if isinstance(v, str) and v.startswith("="):
        v = evaluate(sh, v[1:])
    cache[k] = v
    return v
def evaluate(sh, f):
    def rep(m):
        s2 = m.group(1) or sh
        if m.group(4):
            c0, r0_, c1, r1_ = range_boundaries(f"{m.group(2)}{m.group(3)}:{m.group(4)}{m.group(5)}")
            vals = [val(s2, f"{get_column_letter(c)}{r}") for r in range(r0_, r1_ + 1) for c in range(c0, c1 + 1)]
            return "[" + ",".join(repr(float(x or 0)) for x in vals) + "]"
        return repr(float(val(s2, f"{m.group(2)}{m.group(3)}") or 0))
    return eval(REF.sub(rep, f).replace("SUM(", "sum("))
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                val(ws.title, c.coordinate)
# map sheet titles -> xml parts and inject <v> after every <f>
z = zipfile.ZipFile("cyclopean_ar.xlsx")
wbxml = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
rid = dict(re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbxml))
tgt = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels)) or {}
tgt.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
part = {name: "xl/" + tgt[r].lstrip("/").replace("xl/", "") for name, r in rid.items()}
out = zipfile.ZipFile("cyclopean_ar.tmp", "w", zipfile.ZIP_DEFLATED)
inv = {v: k for k, v in part.items()}
for item in z.infolist():
    data = z.read(item.filename)
    if item.filename in inv:
        sh = inv[item.filename]; x = data.decode()
        def put(m):
            v = cache[(sh, m.group(2))]
            return f'{m.group(1)}{m.group(3)}<v>{v!r}</v></c>'
        x = re.sub(r'(<c r="([A-Z]+\d+)"[^>]*>)(<f>[^<]*</f>)<v\s*/?>(?:</v>)?</c>', put, x)
        x = re.sub(r'(<c r="([A-Z]+\d+)"[^>]*>)(<f>[^<]*</f>)</c>', put, x)
        data = x.encode()
    out.writestr(item, data)
out.close(); z.close(); shutil.move("cyclopean_ar.tmp", "cyclopean_ar.xlsx")
