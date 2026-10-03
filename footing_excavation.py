"""Excavation under the footings and rafts, over their blinding area (ground beams excluded).

Fixed levels (from the bulk excavation level down to the founding level):
  area 1: 1114.85 -> 1113.0, except the two rafts with F.B.L = 1109.5
          (north-west raft 50 cm and west raft 80 cm): 1114.85 -> 1109.5
  area 2: 1120.9  -> 1118.0
  - Footprint: blinding under footings & rafts from footings_concrete.py (FOUNDATIONS.dxf,
    local coordinates), placed on site with the transform matched on the S-BLINDING LINE
    polylines that HATCHED_AREA.dxf also carries (rotation -6.81 deg, residual < 1 mm).
  - Sheet "By footing": every isolated footing F1-F10 with its blinding (A+0.2) x (B+0.2)
    and the area it stands in; the rest of the blinding is rafts / water tank / wall footings.
  - Each blinding cell is assigned to area 1 or area 2 (the nearer one when a footing edge
    sticks slightly out of the hatched outline).
Outputs: footing_excavation.xlsx, footing_excavation.pdf/png
"""
import sys
import numpy as np
from scipy import ndimage as ndi
from shapely import contains_xy, distance, points
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill

TOP = {1: 1114.85, 2: 1120.9}
BOTTOM = {1: 1113.0, 2: 1118.0}
DEEP_BOTTOM = 1109.5                                     # F.B.L of the two hatched rafts
DEEP_PIECES = {"North-west raft 50 cm": (-26.26, 50.54), "West raft 80 cm": (-40.63, 34.33)}
COARSE = 5                                               # 10 cm integration cells

sys.argv = sys.argv[:1]
exec(open("footings_concrete.py").read().split('ws.title = "Footings & rafts"')[0])   # feet, lab, feet_blind
a1, a2 = {}, {}
exec(open("compute_levels.py").read().split("# TIN for")[0], a1)
exec(open("number_area2.py").read().split("tin = ")[0], a2)
area1, area2 = a1["area_poly"], a2["area_poly"]

ROT = np.radians(-6.808781732785986)
T = np.array([353780.83527366, 355922.56974443])
Rm = np.array([[np.cos(ROT), -np.sin(ROT)], [np.sin(ROT), np.cos(ROT)]])

# piece of every blinding pixel = nearest footing/raft piece
_, (ir, ic) = ndi.distance_transform_edt(lab == 0, return_indices=True)
owner = np.where(feet_blind, lab[ir, ic], 0)
deep_ids = {name: px(x, y) for name, (x, y) in DEEP_PIECES.items()}

h, w = (owner.shape[0] // COARSE) * COARSE, (owner.shape[1] // COARSE) * COARSE
blk = lambda a: a[:h, :w].reshape(h // COARSE, COARSE, w // COARSE, COARSE)
frac = blk(feet_blind).mean(axis=(1, 3))
deep_frac = {name: blk(owner == k).mean(axis=(1, 3)) for name, k in deep_ids.items()}
r, c = np.nonzero(frac)
cell = RES * COARSE
lx, ly = X0 + (c + 0.5) * cell, Y1 - (r + 0.5) * cell
sx, sy = (Rm @ np.vstack([lx, ly])) + T[:, None]
pts = points(sx, sy)
zone = np.where(distance(area1, pts) <= distance(area2, pts), 1, 2)
a_all = frac[r, c] * cell * cell
a_deep = {name: f[r, c] * cell * cell for name, f in deep_frac.items()}
a_norm = a_all - sum(a_deep.values())

lines = []   # (zone, item, area, top, bottom)
for z in (1, 2):
    m = zone == z
    lines.append((z, "Footings & rafts", a_norm[m].sum(), TOP[z], BOTTOM[z]))
    for name, a in a_deep.items():
        if a[m].sum() > 0.01:
            lines.append((z, f"{name} (F.B.L {DEEP_BOTTOM})", a[m].sum(), TOP[z], DEEP_BOTTOM))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Footing excavation"
ws.append(["Zone", "Item", "Blinding area (m²)", "From level", "To level", "Depth (m)", "Volume (m³)"])
for cc in ws[1]:
    cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="305496")
tot_a = tot_v = 0
sub = {}
for z, item, a, t, b in lines:
    v = a * (t - b)
    ws.append([f"Area {z}", item, round(a, 2), t, b, round(t - b, 2), round(v, 2)])
    tot_a += a; tot_v += v
    s = sub.setdefault(z, [0, 0]); s[0] += a; s[1] += v
ws.append([])
for z, (a, v) in sub.items():
    ws.append([f"Area {z}", "Subtotal", round(a, 2), "", "", "", round(v, 2)])
ws.append(["TOTAL", "", round(tot_a, 2), "", "", "", round(tot_v, 2)])
for cc in ws[ws.max_row]: cc.font = Font(bold=True)
ws.append([])
ws.append(["Footprint", "blinding under footings & rafts, ground beams excluded (FOUNDATIONS.dxf)"])

# ---- per footing: blinding (A+0.2) x (B+0.2), area from the footing's position on site
EXTRA_FOOTINGS = [("F2", 11.2, 6.05), ("F2", 19.5, 6.05), ("F3", 19.4, 9.4)]   # sit inside raft/tank pieces
foots = [(r[1].split()[1], r[4], r[5]) for r in rows if r[1].startswith("Footing") and "(x2)" not in r[1]
         and r[1] != "Footing F3"] + EXTRA_FOOTINGS
foots.sort(key=lambda f: (-round(f[2] / 3), f[1]))
ws2 = wb.create_sheet("By footing")
ws2.append(["No.", "Footing", "A x B (m)", "Blinding (A+0.2)x(B+0.2) (m²)", "Area", "From", "To", "Depth (m)",
            "Volume (m³)", "Site X", "Site Y"])
for cc in ws2[1]:
    cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="305496")
    cc.alignment = openpyxl.styles.Alignment(wrap_text=True)
f_tot = {1: [0, 0], 2: [0, 0]}
foot_marks = []
for i, (f, x, y) in enumerate(foots, 1):
    A_, B_, _ = SCHEDULE[f]
    ab = (A_ + 2 * PROJECTION) * (B_ + 2 * PROJECTION)
    px_, py_ = Rm @ np.array([x, y]) + T
    p = points(px_, py_)
    z = 1 if distance(area1, p) <= distance(area2, p) else 2
    dep = TOP[z] - BOTTOM[z]
    ws2.append([i, f, f"{A_:.2f} x {B_:.2f}", round(ab, 2), f"Area {z}", TOP[z], BOTTOM[z], round(dep, 2),
                round(ab * dep, 2), round(px_, 2), round(py_, 2)])
    f_tot[z][0] += ab; f_tot[z][1] += ab * dep
    foot_marks.append((i, f, px_, py_, z))
ws2.append([])
for z in (1, 2):
    ws2.append(["", f"Footings in area {z}", "", round(f_tot[z][0], 2), "", "", "", "", round(f_tot[z][1], 2)])
ws2.append(["", "ALL FOOTINGS (16)", "", round(f_tot[1][0] + f_tot[2][0], 2), "", "", "", "",
            round(f_tot[1][1] + f_tot[2][1], 2)])
for cc in ws2[ws2.max_row]: cc.font = Font(bold=True)
ws2.append([])
ws2.append(["", "Rafts, water tank and wall footings (rest of the blinding)"]); ws2[ws2.max_row][1].font = Font(bold=True)
for z, item, a, t, b in lines:
    a_r = a - (f_tot[z][0] if item == "Footings & rafts" else 0)
    ws2.append(["", f"Area {z}: {item.replace('Footings & rafts', 'rafts / tank / walls')}", "", round(a_r, 2), f"Area {z}",
                t, b, round(t - b, 2), round(a_r * (t - b), 2)])
ws2.append([])
ws2.append(["", "TOTAL (footings + rafts)", "", round(tot_a, 2), "", "", "", "", round(tot_v, 2)])
for cc in ws2[ws2.max_row]: cc.font = Font(bold=True)
for col, wd in zip("ABCDEFGHIJK", [6, 34, 13, 16, 9, 9, 9, 10, 12, 13, 13]):
    ws2.column_dimensions[col].width = wd
for col, wd in zip("ABCDEFG", [10, 40, 18, 12, 10, 10, 14]):
    ws.column_dimensions[col].width = wd

wb.save("footing_excavation.xlsx")

fig, ax = plt.subplots(figsize=(16, 15), dpi=110)
ax.plot(*area1.exterior.xy, color="magenta", lw=2, label=f"Area 1: {TOP[1]} → {BOTTOM[1]}")
ax.plot(*area2.exterior.xy, color="#00a0c0", lw=2, label=f"Area 2: {TOP[2]} → {BOTTOM[2]}")
deep_any = sum(deep_frac.values())[r, c] > 0.5
for m, col, lbl in [((zone == 1) & ~deep_any, "#f6b26b", None), ((zone == 2), "#93c47d", None),
                    (deep_any, "#a64d79", f"F.B.L {DEEP_BOTTOM} rafts: {TOP[1]} → {DEEP_BOTTOM}")]:
    ax.scatter(sx[m], sy[m], s=0.6, marker="s", color=col, label=lbl)
for i, f, px_, py_, z in foot_marks:
    ax.text(px_, py_, f"{i}\n{f}", ha="center", va="center", fontsize=7, weight="bold",
            bbox=dict(fc="white", ec="black", lw=.5, alpha=.85, pad=1))
ax.set_aspect("equal"); ax.axis("off"); ax.legend(loc="lower left", fontsize=11, markerscale=12)
ax.set_title(f"Excavation under footings & rafts (blinding area)\n"
             + "   ".join(f"Area {z}: {v:,.1f} m³" for z, (a, v) in sub.items()) + f"   —   Total = {tot_v:,.1f} m³",
             fontsize=14, weight="bold")
fig.savefig("footing_excavation.png", bbox_inches="tight"); fig.savefig("footing_excavation.pdf", bbox_inches="tight")
for l in lines: print(l[0], l[1], round(l[2], 2), l[3], l[4], round(l[2] * (l[3] - l[4]), 2))
print("total", round(tot_a, 2), round(tot_v, 2), {z: [round(x, 2) for x in s] for z, s in sub.items()})
print("footings", {z: [round(v, 2) for v in t] for z, t in f_tot.items()})
for m_ in foot_marks: print(m_[0], m_[1], "area", m_[4])
