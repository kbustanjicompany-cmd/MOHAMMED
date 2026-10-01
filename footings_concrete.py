"""Reinforced concrete and blinding of the footings and rafts only (no ground beams),
from FOUNDATIONS.dxf.

  - The RC footprint of all foundations comes from blinding_area.py (blinding area pulled
    back by its 10 cm projection).
  - Ground beams are the narrow strips: a 1.0 m square opening removes everything
    narrower than 1 m and keeps the footings and rafts.
  - Each remaining piece is identified:
      isolated footings - its size matches the footing schedule (A x B); the schedule
                          area A x B and thickness C are used
      rafts            - the piece holds a "RAFT Thickness" tag; measured area x tag thickness
      water tank       - the piece at the tank (section A-A, 400 mm base) also holds F3
    Two F2 footings sit inside the large raft piece; they are taken out of it at 0.5 m.
Outputs: footings_concrete.xlsx, footings_concrete.pdf/png
"""
import numpy as np
from scipy import ndimage as ndi
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill

exec(open("blinding_area.py").read().split("wb = openpyxl.Workbook()")[0])   # blind, rc, grid

SCHEDULE = {  # footing: (A, B, C) in m, from the FOOTING SCHEDULE
    "F1": (5.10, 5.10, 0.8), "F2": (2.30, 1.90, 0.5), "F3": (3.60, 3.60, 0.6), "F4": (4.00, 4.00, 0.6),
    "F5": (3.80, 3.80, 0.6), "F6": (8.75, 5.00, 0.8), "F7": (9.35, 6.00, 0.8), "F8": (6.50, 6.50, 0.9),
    "F9": (6.00, 6.00, 0.9), "F10": (6.70, 6.70, 0.9)}
RAFT_TAGS = [(-40.63, 34.33, 0.8), (-24.58, 28.32, 0.8), (-26.26, 50.54, 0.5), (-32.73, 11.11, 0.6),
             (-14.47, 54.86, 0.6), (3.76, 54.86, 0.6), (12.1, 37.95, 0.6), (20.68, 31.94, 0.6),
             (22.27, 22.84, 0.6), (11.43, -1.25, 0.6), (-0.88, 7.71, 0.6), (-16.91, 11.25, 0.6),
             (-13.59, 0.15, 0.6), (17.56, 55.8, 0.6)]
TANK_POINT, TANK_BASE = (22.0, 14.0), 0.4      # "400mm REINFORCED CONCRETE FOUNDATION" (section A-A)
F2_IN_RAFT, F3_IN_TANK = 2, 1                  # F2 / F3 labels that fall inside those pieces

k = 2 * int(0.5 / RES) + 1
feet = ndi.maximum_filter(ndi.minimum_filter(rc.astype(np.uint8), size=k), size=k).astype(bool)
lab, n = ndi.label(feet)
px = lambda x, y: lab[int((Y1 - y) / RES), int((x - X0) / RES)]

rows = []   # (piece, element, area, thickness, x, y)
for i, sl in enumerate(ndi.find_objects(lab), 1):
    m = lab[sl] == i
    area = m.sum() * RES * RES
    cy, cx = ndi.center_of_mass(m)
    x, y = X0 + (cx + sl[1].start) * RES, Y1 - (cy + sl[0].start) * RES
    w, h = (sl[1].stop - sl[1].start) * RES, (sl[0].stop - sl[0].start) * RES
    err = {f: abs(max(w, h) - max(a, b)) + abs(min(w, h) - min(a, b)) for f, (a, b, c) in SCHEDULE.items()}
    best = min(err, key=err.get)
    match = [best] if err[best] < 0.3 else []
    rafts = [t for tx, ty, t in RAFT_TAGS if px(tx, ty) == i]
    if match:
        a, b, c = SCHEDULE[match[0]]
        rows.append((i, f"Footing {match[0]}", a * b, c, x, y))
    elif px(*TANK_POINT) == i:
        a, b, c = SCHEDULE["F3"]
        rows.append((i, "Footing F3", a * b, c, x, y))
        rows.append((i, "Water tank base", area - a * b, TANK_BASE, x, y))
    elif rafts:
        t = max(set(rafts), key=rafts.count)
        if len(rafts) > 3:   # the large raft piece holds two F2 footings
            a, b, c = SCHEDULE["F2"]
            rows.append((i, "Footing F2 (x2)", F2_IN_RAFT * a * b, c, x, y))
            area -= F2_IN_RAFT * a * b
        rows.append((i, f"Raft {t * 100:g} cm", area, t, x, y))
    else:
        rows.append((i, "Unidentified", area, 0.0, x, y))

rc_area = sum(r[2] for r in rows)
volume = sum(r[2] * r[3] for r in rows)
P = int(round(PROJECTION / RES)); pdisk = np.hypot(*np.mgrid[-P:P + 1, -P:P + 1]) <= P
feet_blind = ndi.binary_dilation(feet, structure=pdisk) & blind
fb_area = feet_blind.sum() * RES * RES

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Footings & rafts"
ws.append(["Piece", "Element", "RC area (m²)", "Thickness (m)", "RC volume (m³)"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for i, el, a, t, x, y in rows:
    ws.append([i, el, round(a, 2), t, round(a * t, 2)])
ws.append([])
summary = {}
for _, el, a, t, *_ in rows:
    key = el.split(" (")[0] if el.startswith("Footing") else el
    s = summary.setdefault(key, [0, 0, 0]); s[0] += a; s[1] += a * t
    s[2] += 2 if "(x2)" in el else 1
ws.append(["", "Summary by element", "RC area (m²)", "Count", "RC volume (m³)"])
for c in ws[ws.max_row]: c.font = Font(bold=True)
for el in sorted(summary):
    a, v, cnt = summary[el]
    ws.append(["", el, round(a, 2), cnt, round(v, 2)])
ws.append([])
ws.append(["", "TOTAL reinforced concrete (no ground beams)", round(rc_area, 2), "", round(volume, 2)])
ws.append(["", "Blinding under footings & rafts (m²)", round(fb_area, 2), THICKNESS, round(fb_area * THICKNESS, 2)])
iso = [r for r in rows if r[1].startswith("Footing")]
iso_area = sum(r[2] for r in iso); iso_vol = sum(r[2] * r[3] for r in iso)
iso_blind = 0.0
for _, el, a, t, *_ in iso:
    f = el.split()[1]; cnt = 2 if "(x2)" in el else 1
    A, Bd, _c = SCHEDULE[f]
    iso_blind += cnt * (A + 2 * PROJECTION) * (Bd + 2 * PROJECTION)
ws.append([])
ws.append(["", "Isolated footings F1-F10 only: RC area (m²)", round(iso_area, 2), "", round(iso_vol, 2)])
ws.append(["", "Isolated footings F1-F10 only: blinding (A+0.2)x(B+0.2) (m²)", round(iso_blind, 2), THICKNESS,
           round(iso_blind * THICKNESS, 2)])
for r in (ws.max_row - 1, ws.max_row):
    for c in ws[r]: c.font = Font(bold=True)
for r in (ws.max_row - 1, ws.max_row):
    for c in ws[r]: c.font = Font(bold=True)
for col, w in zip("ABCDE", [8, 44, 14, 14, 16]):
    ws.column_dimensions[col].width = w
wb.save("footings_concrete.xlsx")

fig, ax = plt.subplots(figsize=(16, 14), dpi=110)
xs = np.linspace(X0, X1, W); ys = np.linspace(Y1, Y0, H)
ax.contourf(xs, ys, (rc & ~feet).astype(float), levels=[0.5, 1.5], colors=["#cccccc"])
ax.contourf(xs, ys, feet.astype(float), levels=[0.5, 1.5], colors=["#9fb7d9"])
for g in blind_lines: ax.plot(*g.xy, color="red", lw=0.4)
done = set()
for i, el, a, t, x, y in rows:
    if i in done:
        continue
    done.add(i)
    txt = "\n".join(f"{e}: {aa:.1f} m² × {tt:g}" for j, e, aa, tt, *_ in rows if j == i)
    ax.text(x, y, txt, ha="center", va="center", fontsize=6.5, weight="bold",
            bbox=dict(fc="white", ec="none", alpha=0.8, pad=1))
ax.set_xlim(X0, X1); ax.set_ylim(Y0, Y1); ax.set_aspect("equal"); ax.axis("off")
ax.set_title(f"Footings & rafts only (blue) — ground beams excluded (grey)\n"
             f"RC area = {rc_area:,.2f} m²   RC volume = {volume:,.2f} m³   |   blinding = {fb_area:,.2f} m²",
             fontsize=14, weight="bold")
fig.savefig("footings_concrete.png", bbox_inches="tight"); fig.savefig("footings_concrete.pdf", bbox_inches="tight")
print(f"RC area={rc_area:.2f} volume={volume:.2f} blinding={fb_area:.2f}")
print(f"isolated footings: area={iso_area:.2f} volume={iso_vol:.2f} blinding={iso_blind:.2f}")
for el in sorted(summary): print(el, [round(v, 2) for v in summary[el]])
