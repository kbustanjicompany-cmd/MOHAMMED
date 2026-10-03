"""Blinding of ALL foundations except ground beams, drawn by type as a DXF (FOUNDATIONS.dxf
coordinates, for overlay) plus PDF and area table.

Types
  - isolated footings F1-F10 : exact rectangles from footing_blinding_dxf.py
  - continuous (strip) footings : parts of the foundation narrower than STRIP_W
                                  (e.g. the curved perimeter wall footing), > 5 m²
  - rafts / mat              : the wide pieces, by the thickness of their RAFT tag
  - water tank base          : the piece at the tank (section A-A)
Blinding of each type = foundation blinding (footings_concrete.py) assigned to the nearest
foundation piece, then traced to closed polylines (2 cm raster, simplified to 1 cm).
Outputs: foundation_blinding.dxf, foundation_blinding.pdf, foundation_blinding.xlsx
"""
import sys
import numpy as np
from skimage import measure
from shapely.geometry import Polygon, box as sbox
from shapely.ops import unary_union
import ezdxf
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill

STRIP_W = 2.0
exec(open("footing_blinding_dxf.py").read().split("doc = ezdxf.new")[0])     # footings_concrete + isolated rects `out`

iso = unary_union([sbox(*o["blind"]) for o in out])
# narrow parts of the footprint = continuous footings
k = 2 * int(STRIP_W / 2 / RES) + 1
wide = ndi.maximum_filter(ndi.minimum_filter(feet.astype(np.uint8), size=k), size=k).astype(bool)
narrow_lab, nn = ndi.label(feet & ~wide)
sizes = ndi.sum(np.ones_like(narrow_lab), narrow_lab, index=np.arange(1, nn + 1)) * RES * RES
strip_rc = np.isin(narrow_lab, [i + 1 for i, a in enumerate(sizes) if a > 5])

# owner piece of every blinding pixel; class per pixel
_, (ir, ic) = ndi.distance_transform_edt(lab == 0, return_indices=True)
owner = np.where(feet_blind, lab[ir, ic], 0)
P = int(round(PROJECTION / RES)); pdisk = np.hypot(*np.mgrid[-P:P + 1, -P:P + 1]) <= P
strip_bl = ndi.binary_dilation(strip_rc, structure=pdisk) & feet_blind

piece_type = {}
for i, el, a, t, x, y in rows:
    if el.startswith("Raft"):
        piece_type[i] = f"RAFT-{int(round(t * 100))}"
    elif el == "Water tank base":
        piece_type[i] = "WATER-TANK"
classes = {}
for kk, name in piece_type.items():
    classes.setdefault(name, np.zeros_like(feet_blind))
    classes[name] |= (owner == kk) & ~strip_bl
classes["STRIP"] = strip_bl


def trace(mask):
    """mask -> shapely (multi)polygon in plan coordinates."""
    m = np.pad(mask, 1)
    rings = []
    for c in measure.find_contours(m.astype(np.uint8), 0.5):
        if len(c) < 4:
            continue
        xy = [(X0 + (cc - 1) * RES + RES / 2, Y1 - (rr - 1) * RES - RES / 2) for rr, cc in c]
        pg = Polygon(xy).buffer(0)
        if pg.area > 0.05:
            rings.append(pg)
    geom = Polygon()
    for pg in sorted(rings, key=lambda g: -g.area):
        geom = geom.symmetric_difference(pg)
    return geom.simplify(0.01)


def drop_slivers(g, min_area=1.0):
    parts = [p for p in getattr(g, "geoms", [g]) if p.area >= min_area]
    return unary_union(parts) if parts else Polygon()


geoms = {name: drop_slivers(trace(m).difference(iso).buffer(0)) for name, m in classes.items()}
geoms = {n: g for n, g in geoms.items() if g.area > 0.5}
geoms["ISOLATED"] = iso
ORDER = ["ISOLATED", "STRIP", "RAFT-50", "RAFT-60", "RAFT-80", "WATER-TANK"]
NAMES = {"ISOLATED": "Isolated footings F1-F10 (16 No.)", "STRIP": "Continuous (strip) footings",
         "RAFT-50": "Raft / mat 50 cm", "RAFT-60": "Raft / mat 60 cm", "RAFT-80": "Raft / mat 80 cm",
         "WATER-TANK": "Water tank base 40 cm"}
COLORS = {"ISOLATED": (1, "#e06666"), "STRIP": (30, "#f6b26b"), "RAFT-50": (3, "#93c47d"),
          "RAFT-60": (5, "#6fa8dc"), "RAFT-80": (6, "#8e7cc3"), "WATER-TANK": (4, "#76d7ea")}

doc = ezdxf.new("R2018", setup=True); doc.units = ezdxf.units.M
msp = doc.modelspace()
doc.layers.add("BLINDING-LABEL", color=7)
areas = {}
for name in ORDER:
    if name not in geoms:
        continue
    g = geoms[name]
    aci, _ = COLORS[name]
    lay = f"BLINDING-{name}"
    doc.layers.add(lay, color=aci); doc.layers.add(lay + "-HATCH", color=aci)
    polys = list(getattr(g, "geoms", [g]))
    areas[name] = (g.area, len(polys))
    for pg in polys:
        if pg.area < 0.05:
            continue
        ext = list(pg.exterior.coords)[:-1]
        msp.add_lwpolyline(ext, close=True, dxfattribs={"layer": lay})
        for hole in pg.interiors:
            msp.add_lwpolyline(list(hole.coords)[:-1], close=True, dxfattribs={"layer": lay})
        h = msp.add_hatch(color=aci, dxfattribs={"layer": lay + "-HATCH"})
        h.set_pattern_fill("ANSI31", scale=0.05)
        h.paths.add_polyline_path(ext, is_closed=True, flags=ezdxf.const.BOUNDARY_PATH_EXTERNAL)
        for hole in pg.interiors:
            h.paths.add_polyline_path(list(hole.coords)[:-1], is_closed=True)
        if name != "ISOLATED" and pg.area > 3:
            c = pg.representative_point()
            msp.add_mtext(f"{NAMES[name]}\\P{pg.area:.2f} m²", dxfattribs={"layer": "BLINDING-LABEL", "char_height": 0.25}
                          ).set_location((c.x, c.y), attachment_point=5)
for o in out:
    bx0, by0, bx1, by1 = o["blind"]
    msp.add_mtext(f"{o['no']} - {o['f']}\\P{o['area']:.2f} m²", dxfattribs={"layer": "BLINDING-LABEL", "char_height": 0.22}
                  ).set_location(((bx0 + bx1) / 2, (by0 + by1) / 2), attachment_point=5)
total = sum(a for a, n in areas.values())
txt = "BLINDING OF FOUNDATIONS (ground beams excluded) - 100 mm\\P" + "\\P".join(
    f"{NAMES[n]}: {areas[n][0]:.2f} m²" for n in ORDER if n in areas) + f"\\PTOTAL = {total:.2f} m²"
msp.add_mtext(txt, dxfattribs={"layer": "BLINDING-LABEL", "char_height": 0.5}).set_location((-50, 66), attachment_point=7)
doc.saveas("foundation_blinding.dxf")

fig, ax = plt.subplots(figsize=(16.5, 13))
for name in ORDER:
    if name not in geoms:
        continue
    _, col = COLORS[name]
    for pg in getattr(geoms[name], "geoms", [geoms[name]]):
        ax.fill(*pg.exterior.xy, fc=col, ec="black", lw=.5, alpha=.85)
        for hole in pg.interiors:
            ax.fill(*hole.xy, fc="white", ec="black", lw=.5)
    ax.fill([], [], fc=col, ec="black", label=f"{NAMES[name]}: {areas[name][0]:,.2f} m²")
for o in out:
    bx0, by0, bx1, by1 = o["blind"]
    ax.text((bx0 + bx1) / 2, (by0 + by1) / 2, f"{o['no']}\n{o['f']}", ha="center", va="center", fontsize=6.5, weight="bold")
ax.set_aspect("equal"); ax.axis("off")
ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.01), ncol=3, fontsize=11, frameon=False)
ax.set_title(f"Blinding of all foundations (ground beams excluded) — total {total:,.2f} m²", fontsize=14, weight="bold")
fig.savefig("foundation_blinding.pdf", bbox_inches="tight")

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Blinding by type"
ws.append(["Type", "Pieces", "Blinding area (m²)", "Volume @ 100 mm (m³)"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for n in ORDER:
    if n in areas:
        ws.append([NAMES[n], areas[n][1], round(areas[n][0], 2), round(areas[n][0] * .1, 2)])
ws.append(["TOTAL", "", round(total, 2), round(total * .1, 2)])
for c in ws[ws.max_row]: c.font = Font(bold=True)
ws.append([])
ws.append(["Isolated footings", "No.", "Blinding (m²)"])
for o in out:
    ws.append([f"{o['no']} - {o['f']}", 1, round(o["area"], 2)])
for col, w in zip("ABCD", [36, 8, 18, 20]):
    ws.column_dimensions[col].width = w
wb.save("foundation_blinding.xlsx")
for n in ORDER:
    if n in areas: print(n, round(areas[n][0], 2), areas[n][1])
print("total", round(total, 2), "raster feet_blind", round(feet_blind.sum() * RES * RES, 2))
