"""September 2026 monthly works report - translated to Arabic and formatted (monthly_report_sep2026.xlsx)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TITLE = "الأعمال والفعاليات المنجزة خلال شهر أيلول 2026"
HEAD = ["الرقم", "الأعمال المنجزة خلال شهر أيلول", "الأعمال المتوقع إنجازها", "المواد الموردة للموقع", "الفحوصات المخبرية"]
DONE = ["صب خرسانة النظافة لقاعدة الحصيرة المستمرة على المحور A (04–11) بطول 30.80 م",
        "صب الخرسانة المسلحة لقاعدة الحصيرة على المحور A (04–11)",
        "صب خرسانة النظافة لقاعدة الحصيرة على المحور 10 (A–F)",
        "صب الخرسانة المسلحة لقاعدة الحصيرة على المحور 10 (A–F)",
        "صب المرحلة الثانية من الجدار الاستنادي والأعمدة على امتداد المحور A (04–11) حتى منسوب 1120.57"]
NEXT = ["صب الجدار الاستنادي في التسوية الثانية (البدروم 2) وخزان المياه والدرج",
        "الردم خلف الجدار الاستنادي بعد الانتهاء من أعمال العزل",
        "صب القاعدة المستمرة في منطقة الشارع بعد حل مشكلة عمود الكهرباء",
        "البدء بأعمال نظام التدعيم (Shoring)"]
MAT = ["خرسانة مسلحة: 299 م³", "خرسانة عادية: 15 م³", "حديد تسليح مبروم: 173 طن"]
LAB = ["فحص الشد والاستطالة لحديد التسليح بأقطار 12، 14، 16، 18، 20، 25 مم",
       "فحص الشد والاستطالة لحديد التسليح بقطري 10، 32 مم"]

F, NAVY, LIGHT, BAND = "Arial", "1F3864", "D9E2F3", "F5F7FB"
thin = Side(style="thin", color="8EA9DB"); thick = Side(style="medium", color=NAVY)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
rgt = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2, indent=1)

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "تقرير شهر أيلول"
ws.sheet_view.rightToLeft = True; ws.sheet_view.showGridLines = False
for col, w in zip("ABCDEF", [2, 8, 46, 42, 28, 38]):
    ws.column_dimensions[col].width = w

ws.merge_cells("B2:F2"); c = ws["B2"]; c.value = TITLE
c.font = Font(name=F, size=16, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = ctr
ws.row_dimensions[2].height = 38
ws.merge_cells("B3:F3"); c = ws["B3"]; c.value = "التقرير الشهري لسير الأعمال"
c.font = Font(name=F, size=11, italic=True, color=NAVY); c.alignment = ctr; ws.row_dimensions[3].height = 22

hr = 5
for i, t in enumerate(HEAD):
    c = ws.cell(hr, 2 + i, t)
    c.font = Font(name=F, size=12, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F5496")
    c.alignment = ctr; c.border = Border(left=thin, right=thin, top=thick, bottom=thick)
ws.row_dimensions[hr].height = 34

n = max(len(DONE), len(NEXT), len(MAT), len(LAB))
for i in range(n):
    r = hr + 1 + i
    fill = PatternFill("solid", fgColor=BAND) if i % 2 else None
    vals = [i + 1] + [lst[i] if i < len(lst) else None for lst in (DONE, NEXT, MAT, LAB)]
    for j, v in enumerate(vals):
        c = ws.cell(r, 2 + j, v)
        c.font = Font(name=F, size=11, bold=(j == 0), color=NAVY if j == 0 else "000000")
        c.alignment = ctr if j == 0 else rgt
        c.border = Border(left=thin, right=thin, top=thin, bottom=thick if i == n - 1 else thin)
        if j == 0: c.fill = PatternFill("solid", fgColor=LIGHT)
        elif fill: c.fill = fill
    ws.row_dimensions[r].height = 48
# outer frame
last = hr + n
for r in range(hr, last + 1):
    for col, side in ((2, "right"), (6, "left")):
        c = ws.cell(r, col); b = c.border
        c.border = Border(left=thick if side == "left" else b.left, right=thick if side == "right" else b.right, top=b.top, bottom=b.bottom)

s = last + 3
for k, (a, b) in enumerate([("C", "إعداد: ............................"), ("D", "مراجعة: ............................"),
                            ("F", "اعتماد: ............................")]):
    c = ws[f"{a}{s}"]; c.value = b; c.font = Font(name=F, size=11, bold=True, color=NAVY); c.alignment = ctr
    c2 = ws[f"{a}{s + 1}"]; c2.value = "التوقيع: ................  التاريخ: ....../....../2026"
    c2.font = Font(name=F, size=10, color="595959"); c2.alignment = ctr

ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.print_area = f"B2:F{s + 1}"
ws.page_margins.left = ws.page_margins.right = 0.4
wb.save("monthly_report_sep2026.xlsx")
