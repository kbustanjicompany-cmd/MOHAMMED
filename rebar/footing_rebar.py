"""Reinforcement drawings and bar bending schedule (BBS) of the isolated footings F1-F10,
from the FOOTING SCHEDULE of FOUNDATIONS.dxf.

Schedule (A x B x C in mm; D/E bottom mesh, D'/E' top mesh, F starter bars under column).
Assumptions (not given in the drawing - check with the structural engineer):
  - D bars run parallel to A, E bars parallel to B (D is the heavier bar where they differ,
    i.e. along the longer side); same for D'/E'
  - cover 75 mm all round; mesh bars end in a vertical leg (hook) of C - 2 x cover
  - starter bars: L-bars with a 300 mm foot, standing C - cover inside the footing and
    projecting 50 x diameter above it for the lap with the column bars
  - unit weight = d^2 / 162 kg/m
Outputs (in rebar/): footing_rebar.dxf, footing_rebar.pdf, footing_bbs.xlsx
"""
import math
import ezdxf
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import openpyxl
from openpyxl.styles import Font, PatternFill

COVER = 0.075
FOOT, LAP = 0.30, 50
# name: count in plan, A, B, C (m), D, E, D', E' as (dia mm, spacing m) or None, F as (n, dia) or None
FOOTINGS = {
    "F1": (2, 5.10, 5.10, 0.8, (20, .15), (20, .15), None, None, (6, 20)),
    "F2": (3, 2.30, 1.90, 0.5, (12, .15), (12, .15), None, None, (6, 12)),
    "F3": (1, 3.60, 3.60, 0.6, (16, .15), (16, .15), None, None, (6, 16)),
    "F4": (1, 4.00, 4.00, 0.6, (18, .15), (18, .15), None, None, (6, 18)),
    "F5": (2, 3.80, 3.80, 0.6, (16, .15), (16, .15), None, None, (6, 16)),
    "F6": (1, 8.75, 5.00, 0.8, (18, .15), (18, .15), (12, .15), (12, .15), None),
    "F7": (2, 9.35, 6.00, 0.8, (25, .15), (18, .15), (12, .15), (12, .15), None),
    "F8": (2, 6.50, 6.50, 0.9, (25, .15), (25, .15), None, None, (6, 25)),
    "F9": (1, 6.00, 6.00, 0.9, (25, .15), (25, .15), None, None, (6, 25)),
    "F10": (1, 6.70, 6.70, 0.9, (25, .15), (25, .15), None, None, (6, 25)),
}
kgm = lambda d: d * d / 162.0


def bars(name):
    """BBS rows of one footing: (mark, layer, dia, spacing, n, shape, legs, length)."""
    cnt, A, B, C, D, E, Dt, Et, F = FOOTINGS[name]
    hook = C - 2 * COVER
    out = []
    for mark, spec, along, across, layer in (("D", D, A, B, "Bottom"), ("E", E, B, A, "Bottom"),
                                             ("D'", Dt, A, B, "Top"), ("E'", Et, B, A, "Top")):
        if spec:
            d, s = spec
            n = math.floor((across - 2 * COVER) / s) + 1
            straight = along - 2 * COVER
            out.append((mark, layer, d, s, n, "U (hooks both ends)",
                        f"{hook:.2f} + {straight:.2f} + {hook:.2f}", straight + 2 * hook))
    if F:
        n, d = F
        vert = C - COVER + LAP * d / 1000
        out.append(("F", "Starter", d, None, n, "L (foot + vertical)", f"{FOOT:.2f} + {vert:.2f}", FOOT + vert))
    return out


# ---------------- BBS workbook
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BBS footings"
hdr = ["Footing", "No. of footings", "Bar mark", "Layer", "Dia (mm)", "Spacing (m)", "Bars / footing",
       "Shape", "Legs (m)", "Bar length (m)", "Total bars", "Total length (m)", "kg/m", "Weight (kg)"]
ws.append(hdr)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
by_dia = {}
grand = 0.0
for name, row in FOOTINGS.items():
    cnt = row[0]
    for mark, layer, d, s, n, shape, legs, L in bars(name):
        tot_n = n * cnt; tot_L = tot_n * L; wt = tot_L * kgm(d)
        ws.append([name, cnt, mark, layer, d, s, n, shape, legs, round(L, 2), tot_n, round(tot_L, 2),
                   round(kgm(d), 3), round(wt, 1)])
        by_dia[d] = by_dia.get(d, 0) + wt; grand += wt
ws.append([])
ws.append(["Summary by diameter"]); ws[ws.max_row][0].font = Font(bold=True)
for d in sorted(by_dia):
    ws.append([f"Ø{d}", "", "", "", d, "", "", "", "", "", "", "", "", round(by_dia[d], 1)])
ws.append(["TOTAL (kg)", "", "", "", "", "", "", "", "", "", "", "", "", round(grand, 1)])
ws.append(["TOTAL (ton)", "", "", "", "", "", "", "", "", "", "", "", "", round(grand / 1000, 3)])
for r in (ws.max_row - 1, ws.max_row):
    for c in ws[r]: c.font = Font(bold=True)
ws.append([])
for line in __doc__.strip().splitlines()[4:11]:
    ws.append([line.strip()])
for col, w in zip("ABCDEFGHIJKLMN", [9, 9, 8, 9, 8, 10, 11, 20, 20, 13, 10, 14, 7, 12]):
    ws.column_dimensions[col].width = w
wb.save("rebar/footing_bbs.xlsx")

# ---------------- drawings (DXF + PDF), one footing per panel
doc = ezdxf.new("R2018", setup=True); doc.units = ezdxf.units.M
msp = doc.modelspace()
for lay, col in (("CONCRETE", 7), ("BLINDING", 1), ("BAR-BOTTOM", 3), ("BAR-TOP", 6), ("BAR-STARTER", 5),
                 ("TEXT", 2), ("DIM", 4)):
    doc.layers.add(lay, color=col)

pdf = PdfPages("rebar/footing_rebar.pdf")
ox = 0.0
for name, (cnt, A, B, C, D, E, Dt, Et, F) in FOOTINGS.items():
    rows = bars(name)
    fig = plt.figure(figsize=(16.5, 11.7))
    gs = fig.add_gridspec(2, 3, height_ratios=[2.2, 1], width_ratios=[1, 1, 0.9])
    axb, axt, axi = fig.add_subplot(gs[0, 0]), fig.add_subplot(gs[0, 1]), fig.add_subplot(gs[0, 2])
    ax2 = fig.add_subplot(gs[1, :])
    c = COVER
    lab = []
    for mark, layer, d, s_, n, shape, legs, L in rows:
        sp = f"/{s_ * 100:.0f}cm" if s_ else ""
        lab.append(f"{mark:<3}{layer:<8}{n:>3} Ø{d}{sp:<7} L={L:5.2f} m" if mark != "F"
                   else f"F  Starter {n:>3} Ø{d}         L={L:5.2f} m")
    for ax, layer, col in ((axb, "Bottom", "green"), (axt, "Top", "purple")):
        ax.add_patch(plt.Rectangle((-B / 2 - .1, -A / 2 - .1), B + .2, A + .2, fill=False, ec="red", lw=.8, ls="--"))
        ax.add_patch(plt.Rectangle((-B / 2, -A / 2), B, A, fill=False, ec="black", lw=2))
        has = False
        for mark, lay, d, s_, n, *_ in rows:
            if lay != layer:
                continue
            has = True
            lw = 0.25 + d / 40
            if mark in ("D", "D'"):
                for k in range(n):
                    x = -(B / 2 - c) + k * s_
                    ax.plot([x, x], [-(A / 2 - c), A / 2 - c], color=col, lw=lw)
                ax.text(0, A / 2 + .15, f"{mark}: {n} Ø{d}/{s_ * 100:.0f}cm  (∥ A)", ha="center", fontsize=10, color=col, weight="bold")
            else:
                for k in range(n):
                    y = -(A / 2 - c) + k * s_
                    ax.plot([-(B / 2 - c), B / 2 - c], [y, y], color=col, lw=lw)
                ax.text(B / 2 + .15, 0, f"{mark}: {n} Ø{d}/{s_ * 100:.0f}cm  (∥ B)", va="center", rotation=90, fontsize=10, color=col, weight="bold")
        if layer == "Bottom" and F:
            n_, d_ = F
            for (x, y) in [(-.2, -.2), (0, -.2), (.2, -.2), (-.2, .2), (0, .2), (.2, .2)][:n_]:
                ax.add_patch(plt.Circle((x, y), .06, color="blue", zorder=5))
        if not has:
            ax.text(0, 0, "no top mesh", ha="center", va="center", fontsize=12, color="#888888")
        ax.annotate("", xy=(B / 2, -A / 2 - .35), xytext=(-B / 2, -A / 2 - .35), arrowprops=dict(arrowstyle="<->"))
        ax.text(0, -A / 2 - .45, f"B = {B * 1000:.0f}", ha="center", va="top", fontsize=10)
        ax.annotate("", xy=(-B / 2 - .35, A / 2), xytext=(-B / 2 - .35, -A / 2), arrowprops=dict(arrowstyle="<->"))
        ax.text(-B / 2 - .45, 0, f"A = {A * 1000:.0f}", ha="right", va="center", rotation=90, fontsize=10)
        ax.set_aspect("equal"); ax.axis("off")
        ax.set_xlim(-B / 2 - 1.1, B / 2 + .9); ax.set_ylim(-A / 2 - 1.0, A / 2 + .6)
        ax.set_title(f"{layer.upper()} REINFORCEMENT", fontsize=12, weight="bold", color=col)
    axi.axis("off")
    axi.text(0, 1, f"{name}   ({cnt} No.)\n{A * 1000:.0f} x {B * 1000:.0f} x {C * 1000:.0f} mm\n\n" + "\n".join(lab)
             + f"\n\ncover {c * 1000:.0f} mm\nhooks = C - 2 x cover = {(C - 2 * c) * 1000:.0f} mm"
             + (f"\nstarter: foot {FOOT * 1000:.0f} mm, lap {LAP}d above footing" if F else ""),
             va="top", fontsize=11, family="monospace", transform=axi.transAxes)
    # ---- section parallel to A (vertical exaggeration none; own wide panel)
    hook = C - 2 * c
    ax2.add_patch(plt.Rectangle((-A / 2 - .1, -.1), A + .2, .1, fc="#dddddd", ec="red"))
    ax2.add_patch(plt.Rectangle((-A / 2, 0), A, C, fill=False, ec="black", lw=2))
    for mark, layer, d, s_, n, *_ in rows:
        r_ = d / 2000
        if mark == "D":
            y = c + (E[0] / 1000 if E else 0) + r_
            ax2.plot([-A / 2 + c, -A / 2 + c, A / 2 - c, A / 2 - c], [c + hook, y, y, c + hook], color="green", lw=2.2)
        if mark == "E":
            for k in range(math.floor((A - 2 * c) / s_) + 1):
                ax2.add_patch(plt.Circle((-A / 2 + c + k * s_, c + r_), max(r_, .02), color="green"))
        if mark == "D'":
            y = C - c - (Et[0] / 1000 if Et else 0) - r_
            ax2.plot([-A / 2 + c + .05, -A / 2 + c + .05, A / 2 - c - .05, A / 2 - c - .05], [c + .1, y, y, c + .1], color="purple", lw=1.6)
        if mark == "E'":
            for k in range(math.floor((A - 2 * c) / s_) + 1):
                ax2.add_patch(plt.Circle((-A / 2 + c + k * s_, C - c - r_), max(r_, .02), color="purple"))
        if mark == "F":
            n_, d_ = F
            top = C + LAP * d_ / 1000
            for x, sgn in ((-.2, -1), (.2, 1)):
                ax2.plot([x + sgn * FOOT, x, x], [c + .06, c + .06, top], color="blue", lw=2.2)
            ax2.text(.3, top, f"F: {n_} Ø{d_} starter bars", fontsize=10, color="blue", va="top")
    ax2.text(A / 2 + .1, C / 2, f"C = {C * 1000:.0f}", va="center", fontsize=10)
    ax2.text(-A / 2, -.18, "100 mm blinding (projects 100 mm)", fontsize=9, va="top", color="red")
    ax2.set_aspect("equal"); ax2.axis("off")
    ax2.set_xlim(-A / 2 - .3, A / 2 + .9); ax2.set_ylim(-.5, C + (LAP * F[1] / 1000 if F else 0) + .3)
    ax2.set_title("SECTION parallel to A", fontsize=12, weight="bold")
    fig.suptitle("Footing reinforcement from the FOOTING SCHEDULE — assumptions listed in the BBS (check with the structural engineer)",
                 fontsize=10, color="#555555")
    fig.tight_layout(rect=(0, 0, 1, .97))
    pdf.savefig(fig); plt.close(fig)

    # ---- DXF: plan of this footing at (ox, 0)
    cx, cy = ox + B / 2 + 1, A / 2 + 1
    msp.add_lwpolyline([(cx - B / 2, cy - A / 2), (cx + B / 2, cy - A / 2), (cx + B / 2, cy + A / 2), (cx - B / 2, cy + A / 2)],
                       close=True, dxfattribs={"layer": "CONCRETE"})
    msp.add_lwpolyline([(cx - B / 2 - .1, cy - A / 2 - .1), (cx + B / 2 + .1, cy - A / 2 - .1), (cx + B / 2 + .1, cy + A / 2 + .1),
                        (cx - B / 2 - .1, cy + A / 2 + .1)], close=True, dxfattribs={"layer": "BLINDING"})
    for mark, layer, d, s, n, *_ in rows:
        lay = "BAR-BOTTOM" if layer == "Bottom" else ("BAR-TOP" if layer == "Top" else "BAR-STARTER")
        if mark in ("D", "D'"):
            for k in range(n):
                x = cx - (B / 2 - c) + k * s
                msp.add_line((x, cy - (A / 2 - c)), (x, cy + (A / 2 - c)), dxfattribs={"layer": lay})
        elif mark in ("E", "E'"):
            for k in range(n):
                y = cy - (A / 2 - c) + k * s
                msp.add_line((cx - (B / 2 - c), y), (cx + (B / 2 - c), y), dxfattribs={"layer": lay})
    msp.add_text(f"{name} ({cnt} No.) {A * 1000:.0f}x{B * 1000:.0f}x{C * 1000:.0f}", height=0.25,
                 dxfattribs={"layer": "TEXT"}).set_placement((cx - B / 2, cy + A / 2 + .4))
    for k, t in enumerate(lab):
        msp.add_text(t, height=0.15, dxfattribs={"layer": "TEXT"}).set_placement((cx - B / 2, cy - A / 2 - .5 - k * .3))
    ox += B + 4
pdf.close()
doc.saveas("rebar/footing_rebar.dxf")
print(f"footings steel = {grand:,.1f} kg", {d: round(w, 1) for d, w in sorted(by_dia.items())})
