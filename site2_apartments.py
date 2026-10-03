"""Two adjacent apartments inside the SITE2 building with projections (site2_building.py).

Each apartment: 2 bedrooms, living room, kitchen, 2 bathrooms.
Local frame: u along the straight north-west wall from the street (west) end, v = depth from
the north-west wall towards the south-east. Building depth: 8.65 m (u 0-5.25),
9.23 m (u 5.25-10.65), 9.82 m (u 10.65-16.05). The party wall splits the floor area equally.
Room areas are gross (to wall centre-lines); entrances are on the north-west wall, reached
from the street along the 2.5 m side setback.
This is a concept layout - to be developed by the architect (walls, structure, services).
Outputs: site2_apartments.dxf (site coordinates), site2_apartments.pdf, site2_apartments.xlsx
"""
import sys
import numpy as np
import ezdxf
from shapely.geometry import Polygon, box
from shapely import affinity
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Arc
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.argv = sys.argv[:1]
ns = {}
exec(open("site2_building.py").read().split("doc = ezdxf.new")[0], ns)
strips, theta, lot, inner = ns["strips"], ns["proj_theta"], ns["lot"], ns["inner"]
U0 = strips[0][0]; HI = min(s_[2][1] for s_ in strips)   # rotated-frame origin: street end, north-west face
assert all(abs(s_[2][1] - HI) < 1e-3 for s_ in strips)
DEP = [(s_[0] - U0, s_[1] - U0, HI - s_[2][0]) for s_ in strips]   # (u0, u1, depth)
L_TOT = DEP[-1][1]
depth_at = lambda u: next(d for a, b, d in DEP if a - 1e-9 <= u <= b + 1e-9)


def to_site(u, v):
    x, y = U0 + u, HI - v
    t = np.radians(theta)
    return (x * np.cos(t) - y * np.sin(t), x * np.sin(t) + y * np.cos(t))


def step_poly(u0, u1, v0):
    """room from v0 to the (stepped) south-east wall between u0 and u1."""
    pts = [(u0, v0), (u1, v0)]
    cuts = sorted({u0, u1} | {a for a, b, d in DEP if u0 < a < u1})
    back = []
    for a, b in zip(cuts, cuts[1:]):
        d = depth_at((a + b) / 2)
        back += [(a, d), (b, d)]
    return Polygon(pts + list(reversed(back)))


def rect(u0, u1, v0, v1):
    return box(u0, v0, u1, v1)


# party wall where the two apartments get equal areas
def area_upto(p):
    return sum(max(0, min(b, p) - a) * d for a, b, d in DEP)
lo_, hi_ = 0, L_TOT
for _ in range(60):
    mid = (lo_ + hi_) / 2
    lo_, hi_ = (mid, hi_) if area_upto(mid) < area_upto(L_TOT) / 2 else (lo_, mid)
PARTY = round(lo_, 2)

# ---- apartment A (street side, u 0 - PARTY) and B (east, PARTY - end, mirrored)
F = 3.80            # depth of the front band (living / kitchen)
H0, H1 = 3.65, 5.25  # hall / bath 1 band, measured from the outer end wall
VB1, VB2 = 6.20, 5.20  # bath 1 starts at v = 6.20 ; bath 2 ends / bedroom 2 starts at v = 5.20
rooms = []          # (apt, name_en, name_ar, polygon in local frame)
for apt, m in (("A", lambda u: u), ("B", lambda u: L_TOT - u)):
    span = lambda a, b: tuple(sorted((m(a), m(b))))
    party_side = lambda a: tuple(sorted((m(a), PARTY)))
    for n, ar, p in [
            ("LIVING", "صالة", rect(*span(0, 5.0), 0, F)),
            ("KITCHEN", "مطبخ", rect(*party_side(5.0), 0, F)),
            ("HALL", "ممر", rect(*span(H0, H1), F, VB1)),
            ("BEDROOM 1", "غرفة نوم 1", step_poly(*span(0, H0), F)),
            ("BATH 1", "حمام 1", step_poly(*span(H0, H1), VB1)),
            ("BATH 2", "حمام 2", rect(*party_side(H1), F, VB2)),
            ("BEDROOM 2", "غرفة نوم 2", step_poly(*party_side(H1), VB2))]:
        rooms.append((apt, n, ar, p))
W_B = L_TOT - PARTY

footprint = Polygon([(0, 0), (L_TOT, 0)] + [pt for a, b, d in reversed(DEP) for pt in ((b, d), (a, d))])
aptA = footprint.intersection(box(0, 0, PARTY, 20)); aptB = footprint.intersection(box(PARTY, 0, L_TOT, 20))
cover = sum(p.area for *_, p in rooms)
assert abs(cover - footprint.area) < 0.5, (cover, footprint.area)

# doors (hinge point, direction of the opening wall, swing side) and windows, local frame
DOORS = [  # (u, v, wall 'h' horizontal / 'v' vertical, width, swing sign)
    (3.9, 0.0, "h", 1.0, +1), (L_TOT - 4.9, 0.0, "h", 1.0, +1),            # entrances (NW wall)
    (H0, 4.0, "v", 0.9, -1), (L_TOT - H0, 4.0, "v", 0.9, +1),               # bedroom 1
    (H0 + 0.1, VB1, "h", 0.8, +1), (L_TOT - H1 + 0.1, VB1, "h", 0.8, +1),   # bath 1
    (H1, 4.0, "v", 0.8, +1), (L_TOT - H1, 4.0, "v", 0.8, -1),               # bath 2
    (H1, VB2 + 0.1, "v", 0.85, +1), (L_TOT - H1, VB2 + 0.1, "v", 0.85, -1), # bedroom 2
]
WINDOWS = [(0.0, 0.8, 0.0, 3.0), (1.0, 0.0, 3.0, 0.0), (5.6, 0.0, PARTY - 0.6, 0.0),            # A living (street + NW), kitchen
           (0.0, 4.6, 0.0, 7.6), (1.0, depth_at(1), 2.8, depth_at(1)), (5.9, depth_at(6), PARTY - 0.6, depth_at(6)),
           (L_TOT, 0.8, L_TOT, 3.0), (L_TOT - 3.0, 0.0, L_TOT - 1.0, 0.0), (PARTY + 0.5, 0.0, L_TOT - 5.5, 0.0),
           (L_TOT, 4.8, L_TOT, 8.4), (L_TOT - 2.8, depth_at(L_TOT - 1), L_TOT - 1.0, depth_at(L_TOT - 1)),
           (PARTY + 0.5, depth_at(PARTY + 1), 10.5, depth_at(PARTY + 1)),
           (H0 + 0.4, depth_at(4), H1 - 0.4, depth_at(4)), (L_TOT - H1 + 0.4, depth_at(L_TOT - 4), L_TOT - H0 - 0.4, depth_at(L_TOT - 4))]  # bath 1

# ---------------- table
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Apartments"
ws.append(["Apartment", "Room", "الغرفة", "Size (m)", "Area (m²)"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for apt in ("A", "B"):
    for a_, n, ar, p in [r for r in rooms if r[0] == apt]:
        b = p.bounds
        ws.append([apt, n, ar, f"{b[2] - b[0]:.2f} x {b[3] - b[1]:.2f}", round(p.area, 2)])
    tot = (aptA if apt == "A" else aptB).area
    ws.append([apt, "TOTAL (gross)", "المجموع", "", round(tot, 2)])
    for c in ws[ws.max_row]: c.font = Font(bold=True)
    ws.append([])
ws.append(["", "BUILDING", "المبنى", "", round(footprint.area, 2)])
for col, w in zip("ABCDE", [11, 16, 14, 14, 11]):
    ws.column_dimensions[col].width = w
wb.save("site2_apartments.xlsx")

# ---------------- DXF (site coordinates)
doc = ezdxf.new("R2018", setup=True); doc.units = ezdxf.units.M
msp = doc.modelspace()
for name, col in (("LOT", 7), ("SETBACK-LINE", 3), ("A-WALL-EXT", 7), ("A-WALL-INT", 8), ("A-PARTY-WALL", 1),
                  ("A-DOOR", 4), ("A-WINDOW", 5), ("A-ROOM-NAME", 2), ("A-APT-AREA", 6)):
    doc.layers.add(name, color=col)
msp.add_lwpolyline(list(lot.exterior.coords)[:-1], close=True, dxfattribs={"layer": "LOT"})
msp.add_lwpolyline(list(inner.exterior.coords)[:-1], close=True, dxfattribs={"layer": "SETBACK-LINE", "linetype": "DASHED"})
site_poly = lambda p: [to_site(u, v) for u, v in list(p.exterior.coords)[:-1]]
msp.add_lwpolyline(site_poly(footprint), close=True, dxfattribs={"layer": "A-WALL-EXT", "const_width": 0.25})
for _, n, ar, p in rooms:
    msp.add_lwpolyline(site_poly(p), close=True, dxfattribs={"layer": "A-WALL-INT"})
    c = p.representative_point() if n not in ("HALL",) else p.centroid
    msp.add_mtext(f"{n}\\P{p.area:.1f} m²", dxfattribs={"layer": "A-ROOM-NAME", "char_height": 0.22, "rotation": theta}
                  ).set_location(to_site(c.x, c.y), attachment_point=5)
msp.add_line(to_site(PARTY, 0), to_site(PARTY, depth_at(PARTY)), dxfattribs={"layer": "A-PARTY-WALL", "lineweight": 50})
for apt, poly in (("APARTMENT A", aptA), ("APARTMENT B", aptB)):
    c = poly.centroid
    msp.add_mtext(f"{apt}\\P{poly.area:.2f} m²", dxfattribs={"layer": "A-APT-AREA", "char_height": 0.4, "rotation": theta}
                  ).set_location(to_site(c.x, -1.2), attachment_point=5)
for u, v, w, wd, sg in DOORS:
    if w == "h":
        a0, a1 = to_site(u, v), to_site(u + wd, v); leaf = to_site(u, v + sg * wd)
    else:
        a0, a1 = to_site(u, v), to_site(u, v + wd); leaf = to_site(u + sg * wd, v)
    msp.add_line(a0, leaf, dxfattribs={"layer": "A-DOOR"})
    msp.add_line(a0, a1, dxfattribs={"layer": "A-DOOR", "linetype": "DASHED"})
for u0_, v0_, u1_, v1_ in WINDOWS:
    msp.add_line(to_site(u0_, v0_), to_site(u1_, v1_), dxfattribs={"layer": "A-WINDOW", "lineweight": 70})
doc.saveas("site2_apartments.dxf")

# ---------------- PDF (local frame, readable)
fig, ax = plt.subplots(figsize=(16.5, 11.7))
COL = {"LIVING": "#fff2cc", "KITCHEN": "#fce5cd", "HALL": "#eeeeee", "BEDROOM 1": "#d9ead3",
       "BEDROOM 2": "#d9ead3", "BATH 1": "#cfe2f3", "BATH 2": "#cfe2f3"}
for apt, n, ar, p in rooms:
    xs, ys = p.exterior.xy
    ax.fill(xs, [-y for y in ys], fc=COL[n], ec="#444444", lw=1)
    c = p.centroid if n != "BEDROOM 1" else p.representative_point()
    b = p.bounds
    ax.text(c.x, -c.y, f"{n}\n{b[2] - b[0]:.2f} x {b[3] - b[1]:.2f}\n{p.area:.1f} m²", ha="center", va="center",
            fontsize=8.5 if p.area > 6 else 7, weight="bold")
fx, fy = footprint.exterior.xy
ax.plot(fx, [-y for y in fy], color="black", lw=4)
ax.plot([PARTY, PARTY], [0, -depth_at(PARTY)], color="#b00000", lw=4)
for u, v, w, wd, sg in DOORS:
    if w == "h":
        ax.plot([u, u], [-v, -(v + sg * wd)], color="#1155cc", lw=1.5)
        ax.add_patch(Arc((u, -v), 2 * wd, 2 * wd, theta1=-90 if sg > 0 else 0, theta2=0 if sg > 0 else 90, color="#1155cc", lw=.8))
        ax.plot([u, u + wd], [-v, -v], color="white", lw=4.5)
    else:
        ax.plot([u, u + sg * wd], [-v, -v], color="#1155cc", lw=1.5)
        ax.add_patch(Arc((u, -v), 2 * wd, 2 * wd, theta1=(-90 if sg > 0 else 180), theta2=(0 if sg > 0 else 270), color="#1155cc", lw=.8))
        ax.plot([u, u], [-v, -(v + wd)], color="white", lw=4.5)
for u0_, v0_, u1_, v1_ in WINDOWS:
    ax.plot([u0_, u1_], [-v0_, -v1_], color="#3d85c6", lw=5, solid_capstyle="butt")
ax.annotate("ENTRANCE A", (4.4, 0.15), xytext=(4.4, 1.6), ha="center", fontsize=10, color="#1155cc",
            arrowprops=dict(arrowstyle="->", color="#1155cc"))
ax.annotate("ENTRANCE B", (L_TOT - 4.4, 0.15), xytext=(L_TOT - 4.4, 1.6), ha="center", fontsize=10, color="#1155cc",
            arrowprops=dict(arrowstyle="->", color="#1155cc"))
ax.text(-1.2, -4.3, "STREET (west)\n3.0 m setback", rotation=90, va="center", ha="center", fontsize=11, weight="bold")
ax.text(L_TOT / 2, 2.5, "side yard 2.5 m  (access to both entrances)", ha="center", fontsize=10, color="#555555")
ax.text(PARTY / 2, -10.6, f"APARTMENT A — {aptA.area:.2f} m²", ha="center", fontsize=13, weight="bold")
ax.text(PARTY + W_B / 2, -10.6, f"APARTMENT B — {aptB.area:.2f} m²", ha="center", fontsize=13, weight="bold")
ax.set_aspect("equal"); ax.axis("off"); ax.set_xlim(-2, L_TOT + 1); ax.set_ylim(-11.3, 3.2)
ax.set_title("Two adjacent apartments — each: 2 bedrooms, living room, kitchen, 2 bathrooms\n"
             f"building {footprint.area:.2f} m² (concept layout, gross areas)", fontsize=14, weight="bold")
fig.savefig("site2_apartments.pdf", bbox_inches="tight")
print("party", PARTY, "apt A", round(aptA.area, 2), "apt B", round(aptB.area, 2), "rooms total", round(cover, 2), "footprint", round(footprint.area, 2))
for r in rooms: print(r[0], r[1], round(r[3].area, 2), [round(v, 2) for v in r[3].bounds])
