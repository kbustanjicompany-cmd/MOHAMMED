"""Excavation of the outlined area in SITE2.dxf down to FORMATION = 861.5.

  - area      : the closed polyline on layer 0 (colour 11), 323.95 m²
  - levels    : the EL grid (3 m) - point = centre of the small circle, level = Z of the
                "EL:" text; grid corners without a label are interpolated linearly on a TIN
                of all survey points (V-NODE nodes, H: spot heights and EL points)
                (corners outside the survey points take the nearest survey level)
  - grid method: square average = mean of 4 corners; volume = (avg - 861.5) x area of the
                square inside the outline (negative = fill)
  - check     : the same volume integrated directly on the TIN surface (5 cm cells)
Outputs: site2_excavation.xlsx, site2_excavation.pdf
"""
import numpy as np
import ezdxf
from ezdxf import path as zpath
from shapely.geometry import Polygon, box
from shapely import contains_xy
from scipy.interpolate import LinearNDInterpolator, NearestNDInterpolator
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill

FORMATION, STEP = 861.5, 3.0
msp = ezdxf.readfile("SITE2.dxf").modelspace()
area_poly = next(Polygon([(v.x, v.y) for v in zpath.make_path(e).flattening(0.005)])
                 for e in msp.query('LWPOLYLINE[layer=="0"]') if e.closed and e.dxf.color == 11)
circles = [(c.dxf.center.x, c.dxf.center.y) for c in msp.query('CIRCLE[layer=="C-TOPO-TEXT"]')]
labels = []
for t in msp.query('MTEXT[layer=="C-TOPO-TEXT"]'):
    x, y, z = t.dxf.insert
    cx, cy = min(circles, key=lambda c: np.hypot(c[0] - x, c[1] - y))
    labels.append((cx, cy, z))
X0 = min(x for x, _, _ in labels); Y0 = min(y for _, y, _ in labels)
grid = {(round((x - X0) / STEP), round((y - Y0) / STEP)): z for x, y, z in labels}

pts = [tuple(e.dxf.insert) for e in msp.query('INSERT[layer=="V-NODE"]')]
pts += [tuple(t.dxf.insert) for t in msp.query('TEXT[layer=="Heights"]')]
pts += labels
pts = np.array(pts)
tin = LinearNDInterpolator(pts[:, :2], pts[:, 2])
near = NearestNDInterpolator(pts[:, :2], pts[:, 2])   # corners outside the survey hull


def level(i, j):
    if (i, j) in grid:
        return grid[(i, j)], "EL"
    v = float(tin(X0 + i * STEP, Y0 + j * STEP))
    if np.isnan(v):
        return float(near(X0 + i * STEP, Y0 + j * STEP)), "NEAR"
    return v, "TIN"


minx, miny, maxx, maxy = area_poly.bounds
rows = []
for j in range(int(np.floor((maxy - Y0) / STEP)), int(np.floor((miny - Y0) / STEP)) - 1, -1):
    for i in range(int(np.floor((minx - X0) / STEP)), int(np.floor((maxx - X0) / STEP)) + 1):
        x, y = X0 + i * STEP, Y0 + j * STEP
        a = area_poly.intersection(box(x, y, x + STEP, y + STEP)).area
        if a < 0.01:
            continue
        lv = [level(*c) for c in [(i, j + 1), (i + 1, j + 1), (i + 1, j), (i, j)]]
        avg = sum(v for v, _ in lv) / 4
        d = avg - FORMATION
        rows.append(dict(no=len(rows) + 1, x=x, y=y, area=a, h=[v for v, _ in lv],
                         tin=sum(s != "EL" for _, s in lv), near=sum(s == "NEAR" for _, s in lv), avg=avg, d=d,
                         cut=max(d, 0) * a, fill=max(-d, 0) * a))
cut = sum(r["cut"] for r in rows); fill = sum(r["fill"] for r in rows)

# TIN check on 5 cm cells
R = 0.05
gx, gy = np.meshgrid(np.arange(minx + R / 2, maxx, R), np.arange(miny + R / 2, maxy, R))
inside = contains_xy(area_poly, gx, gy)
z = tin(gx[inside], gy[inside])
ok = ~np.isnan(z)
dz = z[ok] - FORMATION
tin_cut = (np.clip(dz, 0, None) * R * R).sum(); tin_fill = (np.clip(-dz, 0, None) * R * R).sum()
tin_cov = ok.sum() * R * R

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Excavation"
ws.append(["Square No.", "Avg level (m)", "Top-left", "Top-right", "Bottom-right", "Bottom-left",
           "Area inside outline (m²)", "Corners interpolated", f"Depth to {FORMATION} (m)", "Cut (m³)", "Fill (m³)"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = openpyxl.styles.Alignment(wrap_text=True)
for r in rows:
    ws.append([r["no"], round(r["avg"], 3), *[round(v, 3) for v in r["h"]], round(r["area"], 2), r["tin"],
               round(r["d"], 3), round(r["cut"], 2), round(r["fill"], 2)])
ws.append([])
for k, v in [("Area of the outline (m²)", round(area_poly.area, 2)), ("Formation level", FORMATION),
             ("TOTAL CUT - grid method (m³)", round(cut, 2)), ("TOTAL FILL - grid method (m³)", round(fill, 2)),
             ("Check: cut on the TIN surface (m³)", round(tin_cut, 2)), ("Check: fill on the TIN surface (m³)", round(tin_fill, 2)),
             ("Area-weighted mean ground level (m)", round(sum(r["avg"] * r["area"] for r in rows) / sum(r["area"] for r in rows), 3))]:
    ws.append([k, v])
    if "TOTAL" in k:
        for c in ws[ws.max_row]: c.font = Font(bold=True)
for col, w in zip("ABCDEFGHIJK", [34, 13, 11, 11, 13, 12, 14, 11, 13, 11, 11]):
    ws.column_dimensions[col].width = w
wb.save("site2_excavation.xlsx")

fig, ax = plt.subplots(figsize=(16, 11))
ax.plot(*area_poly.exterior.xy, color="magenta", lw=2.5, zorder=3)
for r in rows:
    ax.add_patch(plt.Rectangle((r["x"], r["y"]), STEP, STEP, fc="#f4cccc" if r["d"] > 0 else "#cfe2f3",
                               ec="grey", lw=.6, zorder=1))
    cx, cy = r["x"] + STEP / 2, r["y"] + STEP / 2
    ax.text(cx, cy + .8, str(r["no"]), ha="center", fontsize=9, weight="bold", color="#b00000")
    ax.text(cx, cy + .1, f"{r['avg']:.2f}", ha="center", fontsize=8)
    ax.text(cx, cy - .6, f"{r['d']:+.2f} m", ha="center", fontsize=8)
    ax.text(cx, cy - 1.3, f"{r['cut'] if r['d'] > 0 else -r['fill']:.1f} m³", ha="center", fontsize=8, color="#003366")
for (i, j), zz in grid.items():
    ax.plot(X0 + i * STEP, Y0 + j * STEP, "k+", ms=5)
ax.set_aspect("equal"); ax.axis("off")
ax.set_title(f"Excavation to {FORMATION} — area {area_poly.area:.2f} m²\n"
             f"square no. / average level / depth / volume   —   cut = {cut:,.1f} m³, fill = {fill:,.1f} m³ "
             f"(TIN check: cut {tin_cut:,.1f} m³)", fontsize=13, weight="bold")
fig.savefig("site2_excavation.pdf", bbox_inches="tight")
print(f"area={area_poly.area:.2f} squares={len(rows)} cut={cut:.2f} fill={fill:.2f} tin_cut={tin_cut:.2f} tin_fill={tin_fill:.2f} tin_cov={tin_cov:.2f}")
print("near corners", sum(r["near"] for r in rows)); print("avg range", round(min(r['avg'] for r in rows), 2), round(max(r['avg'] for r in rows), 2), "TIN corners", sum(r["tin"] for r in rows))
